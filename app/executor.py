"""Playwright execution engine.

Extracted out of ``qa_agent.py`` so the runner can be driven from three places
with identical behaviour: the dashboard SSE stream, the CI/CD trigger endpoint,
and the scheduler.

What it adds over the original inline runner:

* full step vocabulary — NAVIGATE / KEYPRESS / SELECT / UPLOAD / HOVER / WAIT /
  CHECK / dialogs / downloads plus a much wider VERIFY_* family
* native dialog interception (a ``confirm()`` no longer silently auto-dismisses)
* download capture with row counting for export assertions
* evidence collection feeding :mod:`verdict`, so an expectation is actually
  checked instead of assumed
* optional parallel workers, cancellation, per-step screenshots, Playwright
  tracing, pagination-aware row lookup and visual-regression diffing
"""

from __future__ import annotations

import asyncio
import csv
import json
import os
import re
import shutil
import time
import uuid
from dataclasses import dataclass, field, replace
from pathlib import Path
from typing import Any, Awaitable, Callable

from playwright.async_api import async_playwright

from . import api_seed
from . import config
from . import credentials as credential_store
from .dom_probes import SHARED_PROBES_JS
from . import contracts
from . import coverage
from . import heuristics
from . import learned
from . import pacing
from . import screens
from . import sources as case_sources
from . import storage
from . import rules as rules_engine
from . import verdict as verdict_engine

log = config.get_logger("executor")

Emit = Callable[[dict], Awaitable[None]]

CRASH_MARKERS = (
    "target closed",
    "target page, context or browser has been closed",
    "browser has been closed",
    "connection refused",
    "session closed",
    "websocket",
    "browser closed",
)

ASSERTION_ACTIONS = {
    "VERIFY_STATE", "VERIFY_VISIBLE", "VERIFY_HIDDEN", "VERIFY_DISABLED",
    "VERIFY_ENABLED", "VERIFY_READONLY", "VERIFY_VALUE", "ASSERT_VALUE",
    "VERIFY_NON_ZERO", "VERIFY_TEXT", "VERIFY_ROW_CONTAINS", "VERIFY_TABLE_ROW",
    "VERIFY_ROW_ABSENT", "VERIFY_URL", "VERIFY_COUNT", "VERIFY_DOWNLOAD",
    "VERIFY_MESSAGE",
    # Dependency-chain and API assertions
    "VERIFY_OPTION_PRESENT", "VERIFY_OPTION_ABSENT", "VERIFY_API",
    # Data-quality assertions for grids and audit dialogs
    "VERIFY_COLUMN_NOT_EMPTY", "VERIFY_NO_NOOP_ROWS",
}

# Options inside custom dropdown widgets (PrimeNG, Material, Ant, headless UI).
OPTION_SELECTORS = (
    "option", "[role='option']", ".p-dropdown-item", ".p-multiselect-item",
    ".mat-option", ".ant-select-item-option", "li[data-value]", ".select2-results__option",
    ".dropdown-item", ".choices__item--choice",
)

# Layer 2 of message detection: a message is *text that appeared after an action
# and was not there before*. Installed once per page, it needs no knowledge of the
# application's CSS, which is the whole point — the static list below could not see
# OrangeHRM's `.oxd-input-field-error-message` until someone added it by hand, and
# the next application will have a class nobody has heard of either.
#
# The noise filters matter more than the detection: a grid refreshing, a clock
# ticking or a row count changing all insert text too.
DOM_DIFF_INSTALL_JS = """() => {
    if (window.__qaDiff) { window.__qaDiff.seen.length = 0; return true; }

    const state = { seen: [], ready: false };
    const IGNORE_TAGS = new Set(['SCRIPT', 'STYLE', 'NOSCRIPT', 'TEMPLATE', 'SVG', 'PATH']);
    // Text that changes on its own is not a response to what we did.
    const NOISE = [
        /^[\\d\\s.,:%/-]+$/,                       // numbers, dates, times, counters
        /^\\d{1,2}:\\d{2}(:\\d{2})?\\s*(am|pm)?$/i,   // clocks
        /^(showing|page|rows?|total|of|items?)\\b/i, // pagination furniture
    ];

    const describe = (el) => {
        if (!el || !el.tagName) return '';
        const tag = el.tagName.toLowerCase();
        const cls = (typeof el.className === 'string' ? el.className : '')
            .trim().split(/\\s+/).filter(Boolean).slice(0, 3);
        return cls.length ? `${tag}.${cls.join('.')}` : tag;
    };

    const insideGrid = (el) => !!(el.closest &&
        el.closest('tbody, .p-datatable-tbody, [role="rowgroup"], table'));

    const visible = (el) => {
        try {
            const rect = el.getBoundingClientRect();
            if (rect.width <= 0 || rect.height <= 0) return false;
            const style = getComputedStyle(el);
            return style.visibility !== 'hidden' && style.display !== 'none'
                && parseFloat(style.opacity || '1') > 0.05;
        } catch (e) { return false; }
    };

    const consider = (node) => {
        if (!node) return;
        if (node.nodeType === Node.TEXT_NODE) { node = node.parentElement; }
        if (!node || node.nodeType !== Node.ELEMENT_NODE) return;
        if (IGNORE_TAGS.has(node.tagName)) return;
        if (node.tagName === 'LABEL') return;              // form furniture
        // A node containing controls is a container that was inserted, not the
        // message itself. Without this, opening a dialog whose label reads
        // "Name is required" was classified as a validation error and failed a
        // positive scenario.
        if (node.querySelector && node.querySelector('input, select, textarea, button')) {
            return;
        }
        // A whole dialog or panel being inserted is not itself a message; the
        // message is the leaf that carries the words.
        const text = (node.innerText || node.textContent || '').trim();
        if (!text || text.length > 220) return;
        if (text.split(/\\s+/).length > 30) return;
        if (NOISE.some((re) => re.test(text))) return;
        if (insideGrid(node)) return;
        if (!visible(node)) return;
        // Prefer the innermost element that still holds the whole text.
        let host = node;
        for (let i = 0; i < 4; i++) {
            const child = Array.from(host.children || []).find(
                (c) => ((c.innerText || c.textContent || '').trim() === text));
            if (!child) break;
            host = child;
        }
        const selector = describe(host);
        if (!selector) return;
        if (state.seen.some((s) => s.text === text)) return;
        state.seen.push({ text, selector, source: 'dom-diff' });
        if (state.seen.length > 25) state.seen.shift();
    };

    const observer = new MutationObserver((records) => {
        if (!state.ready) return;
        for (const record of records) {
            record.addedNodes.forEach(consider);
            // An empty element being filled in is an insertion too: Angular
            // renders the error container up front and only sets its text later.
            if (record.type === 'characterData') consider(record.target);
        }
    });
    // An init script runs before the document has a body, so wait for one.
    const attach = () => {
        if (!document.body) { setTimeout(attach, 30); return; }
        observer.observe(document.body, {
            childList: true, subtree: true, characterData: true,
        });
        // Ignore the initial render — only what happens after a step counts.
        setTimeout(() => { state.ready = true; }, 350);
    };
    attach();
    window.__qaDiff = state;
    return true;
}"""

DOM_DIFF_DRAIN_JS = """() => {
    if (!window.__qaDiff) return [];
    const out = window.__qaDiff.seen.slice();
    window.__qaDiff.seen.length = 0;
    return out;
}"""


MESSAGE_SCAN_JS = """(extra) => {
    const selectors = [
        ...(extra || []),
        'small.p-error', 'span.p-error', 'div.p-error', '.invalid-feedback',
        '.error-message', '.field-error', '.p-toast-detail', '.p-toast-summary',
        '.p-toast-message-error', '.p-toast-message-success', '.toast-message',
        '.alert-danger', '.alert-success', '.text-danger', '.text-success',
        '[role="alert"]', '.p-message-text', '.mat-error', '.ant-form-item-explain-error',
        '.help-block', '.validation-message',
        // OrangeHRM / oxd — a real-world miss found while testing a live app:
        // the field error was on screen but invisible to this scanner, which
        // turned a correct application into a false FAIL.
        '.oxd-input-field-error-message', '.oxd-input-group__message',
        '.oxd-alert-content-text', '.oxd-toast-content-text',
        // Bootstrap 5 / generic form frameworks
        '.form-text.text-danger', '.error-text', '.field-validation-error',
        '.Mui-error', '.MuiFormHelperText-root.Mui-error'
    ];
    const out = [];
    for (const sel of selectors) {
        document.querySelectorAll(sel).forEach(el => {
            const rect = el.getBoundingClientRect();
            const visible = rect.width > 0 && rect.height > 0 &&
                            window.getComputedStyle(el).visibility !== 'hidden' &&
                            window.getComputedStyle(el).display !== 'none';
            const text = (el.innerText || el.textContent || '').trim();
            if (visible && text && text.length < 400) {
                out.push({ text, selector: sel + ' ' + (el.className || '') });
            }
        });
    }
    return out;
}"""


class RunCancelled(Exception):
    """Raised inside a scenario when the user stops the run."""


# ──────────────────────────────────────────────────────────────────────
# Run registry (powers Stop, CI/CD polling and the scheduler)
# ──────────────────────────────────────────────────────────────────────
@dataclass
class RunState:
    run_id: str
    project_name: str
    flow_name: str
    status: str = "running"          # running | complete | cancelled | error | interrupted
    started_at: float = field(default_factory=time.time)
    finished_at: float | None = None
    cancelled: bool = False
    total: int = 0
    completed: int = 0
    passed: int = 0
    failed: int = 0
    unverified: int = 0
    crashed: int = 0
    report: dict | None = None
    error: str = ""
    events: list[dict] = field(default_factory=list)

    def snapshot(self) -> dict:
        return {
            "run_id": self.run_id,
            "project_name": self.project_name,
            "flow_name": self.flow_name,
            "status": self.status,
            "started_at": time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(self.started_at)),
            "elapsed_seconds": int((self.finished_at or time.time()) - self.started_at),
            "total": self.total,
            "completed": self.completed,
            "passed": self.passed,
            "failed": self.failed,
            "unverified": self.unverified,
            "crashed": self.crashed,
            "error": self.error,
            # What a CI step should exit with. Anything but a clean, complete run
            # is a failure -- an errored or cancelled run proved nothing.
            "exit_code": 0 if (self.status == "complete" and not self.failed
                               and not self.crashed) else 1,
        }


def persist_run_state(state: "RunState") -> None:
    """Best-effort durable copy of the run's state (see db.run_state)."""
    try:
        from . import db
        db.save_run_state(state.snapshot())
    except Exception as err:  # noqa: BLE001
        log.warning("Could not persist run state %s: %s", state.run_id, err)


RUN_REGISTRY: dict[str, RunState] = {}
_REGISTRY_LIMIT = 200


def register_run(project_name: str, flow_name: str) -> RunState:
    if len(RUN_REGISTRY) > _REGISTRY_LIMIT:
        for key in sorted(RUN_REGISTRY, key=lambda k: RUN_REGISTRY[k].started_at)[:50]:
            if RUN_REGISTRY[key].status != "running":
                RUN_REGISTRY.pop(key, None)
    state = RunState(run_id=uuid.uuid4().hex[:12], project_name=project_name, flow_name=flow_name)
    RUN_REGISTRY[state.run_id] = state
    persist_run_state(state)
    return state


def cancel_run(run_id: str) -> bool:
    state = RUN_REGISTRY.get(run_id)
    if not state or state.status != "running":
        return False
    state.cancelled = True
    return True


def cancel_all_runs() -> int:
    return sum(1 for rid in list(RUN_REGISTRY) if cancel_run(rid))


# ──────────────────────────────────────────────────────────────────────
# Run configuration
# ──────────────────────────────────────────────────────────────────────
@dataclass
class RunConfig:
    project_name: str
    flow_name: str
    num_cases: int = -1
    mode: str = "headless"                 # headed | headless
    test_types: str = ""
    auth_flow: str = ""
    auth_scenario: str = ""
    dependencies: str = ""
    case_ids: str = ""                     # comma separated _original_id values
    workers: int = 1
    verdict_mode: str = "rules"            # off | rules | ai
    step_shots: str = "fail"               # off | fail | all
    tracing: bool = False
    dep_unique_data: bool = False
    data_matrix: str = ""                  # csv filename under flow/data/
    visual_check: bool = False
    env_name: str = ""
    api_seed: str = ""                     # seed template name to pre-run
    seed_once: bool = True                 # honour contract run_once seeding
    auto_complete: bool = True             # let the engine fill required fields it finds
    relationship_checks: bool = False      # append generated dependency scenarios
    apply_rules: bool = True               # enforce project/flow rules in verdicts
    triggered_by: str = "dashboard"
    sources: str = ""                      # "claude,gemini" -- empty means every source
    speed: str = "normal"                  # fast | normal | careful (app/pacing.py)

    @property
    def headless(self) -> bool:
        return self.mode != "headed"


# ──────────────────────────────────────────────────────────────────────
# Step parsing (string DSL + dict steps)
# ──────────────────────────────────────────────────────────────────────
def _strip_quotes(value: str) -> str:
    value = (value or "").strip()
    if len(value) >= 2 and value[0] == value[-1] and value[0] in "\"'`":
        return value[1:-1]
    return value


def parse_step_string(raw: str) -> dict | None:
    """Translate a natural-language step into an executable step dict.

    Kept behaviour-compatible with the original parser and extended with the
    verbs the AI agents already emit (NAVIGATE, KEYPRESS, SELECT, …).
    """
    raw_str = (raw or "").strip()
    if not raw_str:
        return None
    upper = raw_str.upper()

    if upper.startswith(("NAVIGATE TO ", "GOTO ", "GO TO ")):
        for prefix in ("NAVIGATE TO ", "GOTO ", "GO TO "):
            if upper.startswith(prefix):
                return {"action": "NAVIGATE", "selector": raw_str[len(prefix):].strip()}
    if upper.startswith("NAVIGATE "):
        return {"action": "NAVIGATE", "selector": raw_str[9:].strip()}

    # ── Dependency chaining, polling, API and role switching ──────────
    if upper.startswith(("STORE ", "CAPTURE VALUE ", "REMEMBER ")):
        # STORE #custCode AS m7_code   /   REMEMBER #code AS order_code
        for prefix in ("CAPTURE VALUE ", "STORE ", "REMEMBER "):
            if upper.startswith(prefix):
                rest = raw_str[len(prefix):].strip()
                match = re.split(r"\s+as\s+", rest, maxsplit=1, flags=re.IGNORECASE)
                if len(match) == 2:
                    return {"action": "STORE", "selector": match[0].strip(),
                            "value": _strip_quotes(match[1])}
                return {"action": "STORE", "selector": rest, "value": "captured_value"}

    if upper.startswith(("WAIT FOR STATE ", "WAIT UNTIL ", "POLL UNTIL ")):
        for prefix in ("WAIT FOR STATE ", "WAIT UNTIL ", "POLL UNTIL "):
            if upper.startswith(prefix):
                rest = raw_str[len(prefix):].strip()
                match = re.split(r"\s+(?:is|equals|becomes|contains)\s+", rest,
                                 maxsplit=1, flags=re.IGNORECASE)
                if len(match) == 2:
                    return {"action": "WAIT_FOR_STATE", "selector": match[0].strip(),
                            "value": _strip_quotes(match[1])}
                return {"action": "WAIT_FOR_STATE", "selector": rest, "value": ""}

    if upper.startswith(("API GET ", "API POST ", "API PUT ", "API DELETE ", "API_REQUEST ")):
        parts = raw_str.split(None, 2)
        method = parts[1].upper() if len(parts) > 1 and parts[0].upper() == "API" else "GET"
        rest = parts[2] if len(parts) > 2 else (parts[1] if len(parts) > 1 else "")
        return {"action": "API_REQUEST", "method": method, "selector": rest.strip()}

    if upper.startswith(("SWITCH ROLE ", "LOGIN AS ", "SWITCH TO ROLE ")):
        for prefix in ("SWITCH TO ROLE ", "SWITCH ROLE ", "LOGIN AS "):
            if upper.startswith(prefix):
                return {"action": "SWITCH_ROLE", "value": _strip_quotes(raw_str[len(prefix):])}

    if upper.startswith(("VERIFY OPTION ", "ASSERT OPTION ")):
        rest = raw_str[14:].strip()
        absent = re.search(r"\b(absent|missing|not present|excluded|gone)\b", rest, re.IGNORECASE)
        cleaned = re.sub(r"\b(is\s+)?(absent|missing|not present|excluded|gone|present|listed|"
                         r"available|offered)\b", "", rest, flags=re.IGNORECASE)
        match = re.split(r"\s+in\s+", cleaned.strip(), maxsplit=1)
        value = _strip_quotes(match[0].strip(" ,"))
        selector = match[1].strip() if len(match) == 2 else "select"
        return {"action": "VERIFY_OPTION_ABSENT" if absent else "VERIFY_OPTION_PRESENT",
                "selector": selector, "value": value}

    if upper.startswith(("PRESS ", "KEYPRESS ", "KEY ")):
        for prefix in ("KEYPRESS ", "PRESS ", "KEY "):
            if upper.startswith(prefix):
                rest = raw_str[len(prefix):].strip()
                if " on " in rest.lower():
                    idx = rest.lower().find(" on ")
                    return {
                        "action": "KEYPRESS",
                        "value": _strip_quotes(rest[:idx]),
                        "selector": rest[idx + 4:].strip(),
                    }
                return {"action": "KEYPRESS", "selector": "", "value": _strip_quotes(rest)}

    if upper.startswith("SELECT "):
        rest = raw_str[7:].strip()
        for joiner in (" from ", " in ", " on "):
            if joiner in rest.lower():
                idx = rest.lower().find(joiner)
                return {
                    "action": "SELECT",
                    "value": _strip_quotes(rest[:idx]),
                    "selector": rest[idx + len(joiner):].strip(),
                }
        return {"action": "SELECT", "selector": rest, "value": ""}

    if upper.startswith("UPLOAD "):
        rest = raw_str[7:].strip()
        if " to " in rest.lower():
            idx = rest.lower().find(" to ")
            return {
                "action": "UPLOAD",
                "value": _strip_quotes(rest[:idx]),
                "selector": rest[idx + 4:].strip(),
            }
        return {"action": "UPLOAD", "selector": rest, "value": ""}

    if upper.startswith("HOVER "):
        rest = raw_str[6:].strip()
        if rest.lower().startswith("on "):
            rest = rest[3:].strip()
        return {"action": "HOVER", "selector": rest}

    if upper.startswith("WAIT"):
        digits = re.search(r"(\d+)", raw_str)
        return {"action": "WAIT", "selector": "", "value": digits.group(1) if digits else "1000"}

    if upper.startswith(("CHECK ", "TICK ")):
        prefix = 6 if upper.startswith("CHECK ") else 5
        rest = raw_str[prefix:].strip()
        if rest.lower().startswith("on "):
            rest = rest[3:].strip()
        return {"action": "CHECK", "selector": rest}

    if upper.startswith("UNCHECK "):
        rest = raw_str[8:].strip()
        if rest.lower().startswith("on "):
            rest = rest[3:].strip()
        return {"action": "UNCHECK", "selector": rest}

    if upper.startswith("CLEAR "):
        rest = raw_str[6:].strip()
        if rest.lower().startswith("on "):
            rest = rest[3:].strip()
        return {"action": "CLEAR", "selector": rest}

    if upper.startswith("ACCEPT DIALOG"):
        return {"action": "DIALOG", "selector": "", "value": "accept"}
    if upper.startswith("DISMISS DIALOG"):
        return {"action": "DIALOG", "selector": "", "value": "dismiss"}

    if upper.startswith("CLICK ON "):
        return {"action": "CLICK", "selector": raw_str[9:].strip()}
    if upper.startswith("CLICK "):
        return {"action": "CLICK", "selector": raw_str[6:].strip()}

    if upper.startswith(("TYPE ON ", "TYPE ", "ENTER ", "FILL ")):
        if upper.startswith("TYPE ON "):
            prefix_len = 8
        elif upper.startswith("TYPE "):
            prefix_len = 5
        elif upper.startswith("ENTER "):
            prefix_len = 6
        else:
            prefix_len = 5
        rest = raw_str[prefix_len:].strip()
        lowered = rest.lower()
        for joiner in (" with value ", " with ", " into ", " in "):
            if joiner in lowered:
                idx = lowered.find(joiner)
                if joiner in (" into ", " in "):
                    return {
                        "action": "TYPE",
                        "selector": rest[idx + len(joiner):].strip(),
                        "value": _strip_quotes(rest[:idx]),
                    }
                return {
                    "action": "TYPE",
                    "selector": rest[:idx].strip(),
                    "value": _strip_quotes(rest[idx + len(joiner):]),
                }
        return {"action": "TYPE", "selector": rest, "value": ""}

    if upper.startswith(("VERIFY ", "ASSERT ", "CHECK THAT ", "EXPECT ")):
        for prefix in ("VERIFY ", "ASSERT ", "CHECK THAT ", "EXPECT "):
            if upper.startswith(prefix):
                rest = raw_str[len(prefix):].strip()
                break
        if rest.upper().startswith("ON "):
            rest = rest[3:].strip()
        rest_upper = rest.upper()

        if "IS DISABLED" in rest_upper:
            return {"action": "VERIFY_DISABLED",
                    "selector": rest[: rest_upper.find("IS DISABLED")].strip()}
        if "IS ENABLED" in rest_upper:
            return {"action": "VERIFY_ENABLED",
                    "selector": rest[: rest_upper.find("IS ENABLED")].strip()}
        if "IS READONLY" in rest_upper or "IS READ-ONLY" in rest_upper:
            return {"action": "VERIFY_READONLY",
                    "selector": rest[: rest_upper.find("IS READ")].strip()}
        if "IS HIDDEN" in rest_upper or "IS NOT VISIBLE" in rest_upper:
            marker = "IS HIDDEN" if "IS HIDDEN" in rest_upper else "IS NOT VISIBLE"
            return {"action": "VERIFY_HIDDEN",
                    "selector": rest[: rest_upper.find(marker)].strip()}
        if "URL CONTAINS" in rest_upper or "URL IS" in rest_upper:
            marker = "URL CONTAINS" if "URL CONTAINS" in rest_upper else "URL IS"
            idx = rest_upper.find(marker)
            return {"action": "VERIFY_URL", "selector": "",
                    "value": _strip_quotes(rest[idx + len(marker):])}
        if "ROW ABSENT" in rest_upper or "NOT IN TABLE" in rest_upper or \
           "NO ROW" in rest_upper or "ROW REMOVED" in rest_upper:
            match = re.search(r"(ROW\s+ABSENT|NOT\s+IN\s+TABLE|NO\s+ROW|ROW\s+REMOVED)",
                              rest, re.IGNORECASE)
            sel = rest[: match.start()].strip() if match else "table"
            return {"action": "VERIFY_ROW_ABSENT", "selector": sel or "table",
                    "value": _strip_quotes(rest[match.end():]) if match else ""}
        if "HAS ROW" in rest_upper or "ROW CONTAINS" in rest_upper or "CONTAINS ROW" in rest_upper:
            match = re.search(r"(HAS\s+ROW|ROW\s+CONTAINS|CONTAINS\s+ROW)", rest, re.IGNORECASE)
            sel = rest[: match.start()].strip() if match else "table"
            return {"action": "VERIFY_ROW_CONTAINS", "selector": sel or "table",
                    "value": _strip_quotes(rest[match.end():])}
        if "VALUE > 0" in rest_upper or "IS GREATER THAN 0" in rest_upper or \
           "IS NOT ZERO" in rest_upper or "IS NON-ZERO" in rest_upper:
            match = re.search(r"(VALUE\s*>|IS\s+GREATER|IS\s+NOT|IS\s+NON)", rest, re.IGNORECASE)
            return {"action": "VERIFY_NON_ZERO",
                    "selector": rest[: match.start()].strip() if match else rest}
        if "HAS VALUE" in rest_upper or "VALUE IS" in rest_upper or "EQUALS" in rest_upper:
            match = re.search(r"(HAS\s+VALUE|VALUE\s+IS|EQUALS)", rest, re.IGNORECASE)
            return {"action": "VERIFY_VALUE", "selector": rest[: match.start()].strip(),
                    "value": _strip_quotes(rest[match.end():])}
        if "CONTAINS TEXT" in rest_upper or "HAS TEXT" in rest_upper or "SHOWS" in rest_upper:
            match = re.search(r"(CONTAINS\s+TEXT|HAS\s+TEXT|SHOWS)", rest, re.IGNORECASE)
            return {"action": "VERIFY_TEXT", "selector": rest[: match.start()].strip() or "body",
                    "value": _strip_quotes(rest[match.end():])}
        if "MESSAGE" in rest_upper and ("ERROR" in rest_upper or "VALIDATION" in rest_upper):
            return {"action": "VERIFY_MESSAGE", "selector": "", "value": ""}
        if "DOWNLOAD" in rest_upper:
            return {"action": "VERIFY_DOWNLOAD", "selector": "", "value": ""}
        return {"action": "VERIFY_STATE", "selector": rest}

    # Unrecognised free text: preserve the original fallback (treat as a click).
    return {"action": "CLICK", "selector": raw_str}


def parse_steps(steps: list) -> list[dict]:
    parsed: list[dict] = []
    for step in steps or []:
        if isinstance(step, dict):
            normalised = dict(step)
            action = (normalised.get("action") or "").upper().strip()
            if not action:
                action = "TYPE" if normalised.get("value") else "CLICK"
            normalised["action"] = action
            parsed.append(normalised)
        elif isinstance(step, str):
            converted = parse_step_string(step)
            if converted:
                parsed.append(converted)
    return parsed


# ──────────────────────────────────────────────────────────────────────
# Selector helpers
# ──────────────────────────────────────────────────────────────────────
looks_like_dismiss = heuristics.looks_like_dismiss


def candidate_selectors(step: dict, action: str) -> list[str]:
    """Build an ordered candidate list: recorded alternates first, then guesses."""
    selector = (step.get("selector") or "").strip()
    candidates: list[str] = []

    for alt in step.get("selectors") or []:          # recorder-provided alternates
        if isinstance(alt, str) and alt.strip() and alt.strip() not in candidates:
            candidates.append(alt.strip())
    if selector and selector not in candidates:
        candidates.insert(0, selector)

    clean_id = selector.lstrip("#").strip()
    if clean_id and not any(ch in clean_id for ch in "[]()>:. "):
        if action in ("TYPE", "CLEAR", "SELECT", "UPLOAD"):
            candidates += [
                f'input#{clean_id}', f'input[name="{clean_id}"]',
                f'input[formcontrolname="{clean_id}"]', f'input[placeholder*="{clean_id}"]',
                f'textarea#{clean_id}', f'textarea[name="{clean_id}"]',
                f'select#{clean_id}', f'select[name="{clean_id}"]',
                f'[data-testid="{clean_id}"]', f'[aria-label="{clean_id}"]',
            ]
        else:
            candidates += [
                f'[data-testid="{clean_id}"]', f'[aria-label="{clean_id}"]',
                f'button[name="{clean_id}"]',
            ]

    # A label the application spells differently ("Update" for "Save", "Cancel"
    # for "No", an icon button for "Export") is the same control, so offer its
    # other spellings before the step is allowed to fail.
    if action in ("CLICK", "DOUBLE_CLICK"):
        candidates += heuristics.synonym_selectors(selector)

    if action == "CLICK" and looks_like_dismiss(selector):
        candidates += [
                'button:has-text("Cancel")', 'button:has-text("Close")',
                'button:has-text("No")', 'button:has-text("Dismiss")',
                '.p-dialog-footer button:has-text("Cancel")',
                'button.p-dialog-header-close', '.p-dialog-header-close',
                '.modal-footer button:has-text("Cancel")', '.modal-header .btn-close',
                'button[aria-label="Close"]', 'button[type="reset"]',
                'svg.close-icon', 'i.pi-times',
            ]

    seen: set[str] = set()
    ordered: list[str] = []
    for cand in candidates:
        if cand and cand not in seen:
            seen.add(cand)
            ordered.append(cand)
    return ordered


def is_destructive(step: dict) -> bool:
    blob = f"{step.get('selector','')} {step.get('text','')} {step.get('value','')}".lower()
    return any(word in blob for word in config.DESTRUCTIVE_KEYWORDS)


def humanize_error(error: Exception | str, action: str = "", selector: str = "") -> str:
    """Turn a Playwright stack message into something a manual tester can act on."""
    raw = str(error)
    low = raw.lower()
    target = f" '{selector}'" if selector else ""

    # An assertion has already produced a precise, human-written explanation
    # ("Row 'Dinesh' was not found in the table"). This must be checked FIRST:
    # the generic rules below would otherwise rewrite it into "the element was
    # not found… the field name changed", which sends the reader hunting for a
    # selector problem that does not exist.
    if "validation failed:" in low:
        return raw.split("Validation failed:", 1)[-1].strip() or raw

    # Our own diagnostics are already phrased for a tester — do not rewrite them.
    if "the control is disabled" in low or "no element on the page matches" in low:
        return raw.strip()[:400]

    if "timeout" in low and "waiting for" in low:
        return (
            f"Could not interact with the {action or 'element'}{target} in time. "
            "It may be hidden behind a dialog, disabled, or the page may have "
            "changed since the flow was recorded."
        )
    if "strict mode violation" in low:
        return (
            f"The locator{target} matched several elements, so the step was "
            "ambiguous. Make the selector more specific (add a row or table scope)."
        )
    if "not found" in low or "no element" in low or "count = 0" in low:
        return (
            f"The {action or 'element'}{target} was not found on the page. The "
            "screen may not have loaded yet, or the field name changed."
        )
    if "not visible" in low or "not enabled" in low or "element is not" in low:
        return (
            f"The {action or 'element'}{target} exists but is not usable "
            "(hidden or disabled) at this point in the flow."
        )
    if any(marker in low for marker in CRASH_MARKERS):
        return "The browser closed unexpectedly during this scenario."
    if "net::" in low or "err_" in low.upper().lower():
        return f"The page failed to load ({raw.strip()[:160]}). Check the target URL and VPN/network access."
    if "validation failed:" in low:
        return raw.split("Validation failed:", 1)[-1].strip() or raw
    return raw.strip()[:400]


def infer_dialog_policy(scenario: dict) -> str:
    """Decide whether native confirm() dialogs should be accepted or dismissed."""
    blob = f"{scenario.get('scenario_name','')} {scenario.get('expected_result','')}".lower()
    if re.search(r"\bcancel\b|\bdismiss\b|\bdo not (delete|remove|save)\b|\babort\b", blob) and \
       re.search(r"\bconfirm|dialog|delete|remove\b", blob):
        return "dismiss"
    return "accept"


# ──────────────────────────────────────────────────────────────────────
# Data helpers
# ──────────────────────────────────────────────────────────────────────
def load_data_matrix(flow_path: Path, filename: str) -> list[dict]:
    """Load a CSV data matrix for data-driven execution."""
    if not filename:
        return []
    csv_path = flow_path / "data" / filename
    if not csv_path.exists():
        csv_path = flow_path / filename
    if not csv_path.exists():
        return []
    with open(csv_path, "r", encoding="utf-8-sig", newline="") as fh:
        return [dict(row) for row in csv.DictReader(fh)]


def unique_token() -> str:
    return time.strftime("%H%M%S") + uuid.uuid4().hex[:3]


CREDENTIAL_FIELD_MARKERS = (
    "user", "login", "email", "pass", "pwd", "otp", "captcha", "token", "mobile", "phone",
)


def make_dep_value_unique(selector: str, value: str, token: str) -> str:
    """Suffix prerequisite-flow input values so replays don't hit duplicate keys.

    A ``{{variable}}`` is left alone. Reusing a captured value is always
    deliberate — it is how a duplicate test submits the same value twice — so
    tokenising it silently produced a *different* value and the duplicate never
    happened. The scenario then "passed" having proven nothing.
    """
    if not value or not value.strip():
        return value
    if "{{" in value:
        return value
    low = selector.lower()
    if any(marker in low for marker in CREDENTIAL_FIELD_MARKERS):
        return value
    if re.fullmatch(r"[\d\s./:-]+", value.strip()):     # numbers, dates, times
        return value
    if len(value) > 40:
        return value
    return f"{value}{token}"


# ──────────────────────────────────────────────────────────────────────
# Evidence collection
# ──────────────────────────────────────────────────────────────────────
EMPTY_STATE_PHRASES = (
    "no records found", "no record found", "no records", "no data found",
    "no data available", "no data", "no results found", "no results",
    "no matching records", "no matching", "nothing to show", "nothing found",
    "no rows", "no entries", "empty", "0 records",
)


def looks_like_empty_state(text: str) -> bool:
    lowered = (text or "").strip().lower()
    return bool(lowered) and any(p in lowered for p in EMPTY_STATE_PHRASES)


def find_empty_state(text: str) -> str:
    """The empty-grid wording this application actually uses, if it is showing one."""
    lowered = (text or "").lower()
    for phrase in EMPTY_STATE_PHRASES:
        index = lowered.find(phrase)
        if index >= 0:
            return text[index:index + len(phrase)]
    return ""


def _same_input(typed: str, actual: str) -> bool:
    """Did the field keep what was typed? Trimming and formatting don't count."""
    left = (typed or "").strip()
    right = (actual or "").strip()
    if left == right:
        return True
    # Numeric widgets reformat legitimately: 2344 -> 2,344.00 is still 2344.
    strip_format = lambda s: re.sub(r"[,\s]", "", s).rstrip("0").rstrip(".")
    if strip_format(left) and strip_format(left) == strip_format(right):
        return True
    return False


KEYSTROKE_WIDGET_JS = """el => {
    if (!el || el.tagName !== 'INPUT' || el.type === 'number') return false;
    if (el.closest('p-inputnumber, .p-inputnumber')) return true;
    if (el.getAttribute('role') === 'spinbutton') return true;
    return false;
}"""


async def is_keystroke_widget(loc) -> bool:
    """A masked numeric input that only reacts to real key events."""
    try:
        return bool(await loc.evaluate(KEYSTROKE_WIDGET_JS))
    except Exception:
        return False


async def measure_video_size(browser) -> dict:
    """The recording size for windows of this browser: its page in device pixels."""
    try:
        context = await browser.new_context(no_viewport=True)
        try:
            page = await context.new_page()
            w, h = await page.evaluate(
                "() => [innerWidth * devicePixelRatio, innerHeight * devicePixelRatio]")
            return screens.video_size(w, h)
        finally:
            await context.close()
    except Exception:
        return dict(screens.DEFAULT_VIDEO)


FLASH_OUTLINE_JS = """el => {
    const before = el.style.outline, offset = el.style.outlineOffset;
    el.style.outline = '2px solid #f59e0b'; el.style.outlineOffset = '2px';
    setTimeout(() => { el.style.outline = before; el.style.outlineOffset = offset; }, 700);
}"""


async def note_input_rejection(page, loc, step: dict, typed: str, evidence) -> None:
    """Record when a widget refuses the value instead of storing it.

    A numeric field that silently drops letters, or a maxlength that truncates,
    *is* the application enforcing its rule — at the input layer rather than on
    submit. Without this the run reached the end with a clean form and reported
    "the invalid input appears to have been accepted", which is the opposite of
    what happened. The verdict engine reads this to resolve such expectations.
    """
    if not typed or not str(typed).strip():
        return
    try:
        actual = await loc.input_value(timeout=1500)
    except Exception:
        try:
            actual = await loc.evaluate("el => el.value ?? el.textContent ?? ''")
        except Exception:
            return
    if _same_input(str(typed), str(actual)):
        return
    field_name = (step.get("selector") or "").strip()
    if str(actual).strip() == "":
        how = "refused the value outright — the field is still empty"
    elif str(typed).startswith(str(actual)) or str(actual) in str(typed):
        how = f"kept only {str(actual)[:60]!r}"
    else:
        how = f"changed it to {str(actual)[:60]!r}"
    evidence.input_rejected.append(
        {"selector": field_name, "typed": str(typed)[:120], "kept": str(actual)[:120],
         "detail": f"{field_name or 'the field'} {how}"}
    )
    evidence.notes.append(
        f"Input filtered at the keyboard layer: {field_name or 'field'} was given "
        f"{str(typed)[:40]!r} and {how}."
    )


def note_created_record(pool: dict, messages: list[dict]) -> None:
    """After a success message, the value typed first becomes the created record.

    Later assertions should verify that record rather than a literal copied from
    the original recording, which this run never created.
    """
    if pool.get("created_record"):
        return
    if any(m.get("kind") == "success" for m in messages):
        identity = pool.get("first_typed_value") or pool.get("last_typed_value")
        if identity:
            pool["created_record"] = identity


async def install_message_observer(page) -> None:
    """Watch for text that appears, on this page and on every later navigation.

    ``add_init_script`` covers the navigations — a single-page app replacing its
    body, or a real page load — because an observer installed once is lost the
    moment the document changes.
    """
    if not config.DOM_DIFF_MESSAGES:
        return
    try:
        # Self-invoking: a bare "() => {…}" passed as a script body only declares
        # a function and throws it away, so the observer never started and every
        # navigation lost it.
        await page.add_init_script(f"({DOM_DIFF_INSTALL_JS})();")
    except Exception:
        pass
    try:
        await page.evaluate(DOM_DIFF_INSTALL_JS)      # the document already open
    except Exception:
        pass


async def scan_messages(page, evidence: verdict_engine.Evidence,
                        extra_selectors: list[str] | None = None,
                        project_name: str = "") -> list[str]:
    """Harvest visible inline/toast messages before they fade away.

    ``extra_selectors`` are the message surfaces this application has already been
    seen to use, persisted per project by :func:`learned.remember_message_selector`.
    Passing them in is what closes the loop: the platform was recording where a
    message appeared and then scanning without that knowledge on the next run,
    so an application whose error class is in no framework list stayed invisible
    and a correct app was reported as a false FAIL.
    """
    fresh: list[str] = []
    found: list[dict] = []
    try:
        found = await page.evaluate(MESSAGE_SCAN_JS, list(extra_selectors or [])) or []
    except Exception:
        found = []

    # Layer 2: whatever appeared since the last drain, whether or not its CSS is
    # in any list we hold. This is what makes an unknown framework readable.
    if config.DOM_DIFF_MESSAGES:
        try:
            found = list(found) + (await page.evaluate(DOM_DIFF_DRAIN_JS) or [])
        except Exception:
            pass

    for item in found:
        text = (item.get("text") or "").strip()
        if not text:
            continue
        selector = item.get("selector", "")
        before = len(evidence.messages)
        evidence.add_message(text, selector)
        if len(evidence.messages) <= before:
            continue
        fresh.append(text)
        # Only a diff-discovered surface is worth learning — the seed scan can
        # only ever report a selector we already had.
        if item.get("source") == "dom-diff" and project_name:
            kind = verdict_engine.classify_message(text, selector)
            if kind in ("error", "success"):
                learned.remember_message_selector(project_name, selector)
    return fresh


async def wait_for_settle(page, timeout: int = 4000, quiet_ms: int = 300) -> None:
    """Wait while the page is busy, at most ``timeout`` ms (see app/pacing.py)."""
    try:
        await page.wait_for_load_state("domcontentloaded", timeout=timeout)
    except Exception:
        pass
    await pacing.settle(page, timeout, quiet_ms)


# ──────────────────────────────────────────────────────────────────────
# Table verification (pagination + search aware)
# ──────────────────────────────────────────────────────────────────────
ROW_TABLE_SELECTORS = (
    "table tbody", ".p-datatable-tbody", ".table tbody", "[role='grid']", "tbody",
    ".ag-center-cols-container", ".mat-mdc-table",
)
SEARCH_INPUT_SELECTORS = (
    "input[type='search']", ".p-input-icon-left input", "input[placeholder*='Search' i]",
    "input[placeholder*='search' i]", "input[aria-label*='Search' i]",
    ".p-datatable-header input", "#globalFilter", "input.form-control[placeholder*='earch']",
)
NEXT_PAGE_SELECTORS = (
    "button.p-paginator-next:not(.p-disabled)", ".p-paginator-next:not(.p-disabled)",
    "a[aria-label='Next']", "button[aria-label='Next Page']", "li.next:not(.disabled) a",
    ".pagination .page-item:not(.disabled) a[rel='next']", "button:has-text('Next')",
)


async def _page_contains_token(page, selector: str, token: str) -> bool:
    needle = token.lower()
    for table in (selector, *ROW_TABLE_SELECTORS):
        if not table:
            continue
        try:
            loc = page.locator(table).first
            if await loc.count() > 0:
                text = (await loc.inner_text() or "").lower()
                if needle in text:
                    return True
        except Exception:
            continue
    try:
        row = page.locator(f"tr:has-text('{token}')").first
        if await row.count() > 0 and await row.is_visible():
            return True
    except Exception:
        pass
    try:
        return bool(await page.evaluate(
            """(needle) => {
                const cells = document.querySelectorAll(
                    'table td, .p-datatable-tbody td, [role="gridcell"], .ag-cell');
                const n = needle.toLowerCase();
                for (const c of cells) {
                    const txt = (c.innerText || '').toLowerCase();
                    const title = (c.getAttribute('title') || '').toLowerCase();
                    const aria = (c.getAttribute('aria-label') || '').toLowerCase();
                    if (txt.includes(n) || title.includes(n) || aria.includes(n)) return true;
                }
                return false;
            }""",
            token,
        ))
    except Exception:
        return False


async def find_row(page, selector: str, expected: str, *, max_pages: int = 6,
                   use_search: bool = True) -> tuple[bool, str]:
    """Look for a table row across search box and pagination, not just page 1."""
    tokens = [t for t in re.split(r"\s+", (expected or "").strip()) if len(t) > 2]
    token = tokens[0] if tokens else (expected or "").strip()
    if not token:
        return False, "no expected value supplied"

    if await _page_contains_token(page, selector, expected) or \
       await _page_contains_token(page, selector, token):
        return True, "found on the current page"

    if use_search:
        for search_sel in SEARCH_INPUT_SELECTORS:
            try:
                box = page.locator(search_sel).first
                if await box.count() > 0 and await box.is_visible():
                    await box.fill(token, timeout=3000)
                    await box.press("Enter")
                    await wait_for_settle(page, 3000)
                    if await _page_contains_token(page, selector, token):
                        return True, f"found after searching for {token!r}"
                    break
            except Exception:
                continue

    for page_no in range(2, max_pages + 1):
        advanced = False
        for next_sel in NEXT_PAGE_SELECTORS:
            try:
                nxt = page.locator(next_sel).first
                if await nxt.count() > 0 and await nxt.is_visible() and await nxt.is_enabled():
                    await nxt.click(timeout=3000)
                    await wait_for_settle(page, 3000)
                    advanced = True
                    break
            except Exception:
                continue
        if not advanced:
            break
        if await _page_contains_token(page, selector, token):
            return True, f"found on table page {page_no}"

    return False, f"{expected!r} (token {token!r}) not found on any inspected table page"


# ──────────────────────────────────────────────────────────────────────
# Visual regression (optional — needs Pillow)
# ──────────────────────────────────────────────────────────────────────
def compare_visual(baseline: Path, current: Path, diff_out: Path) -> dict:
    try:
        from PIL import Image, ImageChops
    except Exception:
        return {"supported": False,
                "note": "Visual diffing needs Pillow (pip install pillow)."}
    try:
        with Image.open(baseline) as base_img, Image.open(current) as cur_img:
            base_rgb = base_img.convert("RGB")
            cur_rgb = cur_img.convert("RGB")
            if base_rgb.size != cur_rgb.size:
                cur_rgb = cur_rgb.resize(base_rgb.size)
            diff = ImageChops.difference(base_rgb, cur_rgb)
            bbox = diff.getbbox()
            histogram = diff.convert("L").histogram()
            changed = sum(count for value, count in enumerate(histogram) if value > 24)
            total = base_rgb.size[0] * base_rgb.size[1]
            pct = round((changed / total) * 100, 3) if total else 0.0
            if bbox:
                diff.save(diff_out)
            return {"supported": True, "diff_percent": pct,
                    "diff_image": str(diff_out) if bbox else "", "changed_pixels": changed}
    except Exception as err:
        return {"supported": False, "note": f"Visual diff failed: {err}"}


# ──────────────────────────────────────────────────────────────────────
# Single-step execution
# ──────────────────────────────────────────────────────────────────────
async def execute_step(page, step: dict, ctx: dict) -> dict:
    """Run one step. Returns a dict describing what happened.

    ``ctx`` carries: evidence, runtime_pool, headless, log (async callable),
    dialog_policy (mutable dict), downloads list.
    """
    action = (step.get("action") or "").upper()
    # Placeholders are resolved in the selector too: targeting a row by the value
    # a previous module produced — tr:has-text('{{order_code}}') — is the whole
    # point of a dependency chain.
    selector = heuristics.resolve_dynamic_variables(
        (step.get("selector") or "").strip(), ctx["runtime_pool"])
    raw_value = step.get("value", "")
    value = heuristics.resolve_dynamic_variables(str(raw_value or ""), ctx["runtime_pool"])
    evidence: verdict_engine.Evidence = ctx["evidence"]
    emit_log = ctx["log"]
    headless = ctx["headless"]
    if selector != (step.get("selector") or "").strip():
        step = dict(step, selector=selector)
    result = {"action": action, "selector": selector, "healed": False}

    async def highlight(loc) -> None:
        # Not loc.highlight(): that pins the selector text over the page, which
        # covered the very fields a tester is watching. A plain outline instead.
        if not headless:
            try:
                await loc.evaluate(FLASH_OUTLINE_JS)
            except Exception:
                pass

    async def first_usable(cands: list[str], *, require_visible: bool = True):
        """First candidate that resolves — inside the open dialog when there is one.

        Without the scoping step this used ``.first()`` on the whole page, which
        takes whatever appears first in the DOM. With a dialog open that is
        usually an element *behind* it: a value meant for the dialog's "Days"
        field went into a look-alike on the page underneath, and a Status
        dropdown selector matched the pagination dropdown in the footer.
        """
        surface = await mark_active_surface(page)

        async def attempt(sel: str, *, wait: bool):
            """Resolve a selector to a *usable* element, not merely the first one.

            ``.first`` alone made a whole class of runs fail on applications that
            keep closed dialogs in the DOM: `p-dialog` matched the hidden one that
            was rendered first, and the step was reported as "exists but is not
            visible" while the dialog the user could actually see sat two nodes
            later. So scan the matches and take the first that is visible.
            """
            root = page.locator(sel)
            if wait and require_visible:
                try:
                    await root.first.wait_for(state="visible", timeout=4000)
                    return root.first
                except Exception:
                    pass                      # fall through to the match scan
            try:
                total = await root.count()
            except Exception:
                return None
            if total == 0:
                return None
            if not require_visible:
                return root.first
            for index in range(min(total, 12)):
                candidate_loc = root.nth(index)
                try:
                    if await candidate_loc.is_visible():
                        return candidate_loc
                except Exception:
                    continue
            return None

        async def resolve(sel: str, *, wait: bool):
            try:
                loc = await attempt(sel, wait=wait)
                if loc is None and wait:
                    loc = await attempt(sel, wait=False)
                return loc
            except Exception:
                return None

        # Two passes, and the order matters. Every candidate is tried inside the
        # open dialog before *any* of them is tried on the page: a recorded
        # `button:has-text("Save")` would otherwise match a Save button on the
        # page behind the dialog, in preference to the dialog's own button that
        # happens to be labelled "Update". Whatever the dialog offers is what the
        # user is looking at, so it wins outright.
        if surface:
            for position, cand in enumerate(cands):
                scoped = scope_to_surface(cand)
                if not scoped:
                    continue
                loc = await resolve(scoped, wait=position == 0)
                if loc is not None:
                    return loc, scoped

        for position, cand in enumerate(cands):
            if not cand:
                continue
            loc = await resolve(cand, wait=position == 0 and not surface)
            if loc is not None:
                return loc, cand
        return None, ""

    # ── Navigation ────────────────────────────────────────────────────
    if action == "NAVIGATE":
        target = value or selector
        target = heuristics.resolve_dynamic_variables(target, ctx["runtime_pool"])
        if not target or target.lower() in ("window", "page", "browser"):
            result["skipped"] = "no URL supplied"
            return result
        if not target.startswith(("http://", "https://")):
            base = ctx.get("base_url") or ""
            target = base.rstrip("/") + "/" + target.lstrip("/") if base else target
        await page.goto(target, wait_until="domcontentloaded", timeout=60000)
        await wait_for_settle(page)
        result["navigated_to"] = target
        return result

    if action in ("WAIT", "SLEEP"):
        try:
            millis = max(0, min(int(float(value or 1000)), 30000))
        except Exception:
            millis = 1000
        pace = ctx.get("pace") or pacing.pace_for("")
        if pace.wait_steps_fixed:
            await page.wait_for_timeout(millis)
        else:
            # "Up to": stop as soon as the page is quiet. A WAIT written for a
            # slow screen no longer costs its full time on a fast one.
            await pacing.settle(page, millis, pace.quiet_ms)
        return result

    if action in ("WAIT_FOR", "WAIT_FOR_ELEMENT"):
        loc, used = await first_usable(candidate_selectors(step, "CLICK"))
        if not loc:
            raise Exception(f"Validation failed: Element '{selector}' never became visible.")
        result["selector_used"] = used
        return result

    if action == "DIALOG":
        ctx["dialog_policy"]["mode"] = "dismiss" if "dismiss" in value.lower() else "accept"
        await emit_log(f"🗨️ Native dialogs will now be {ctx['dialog_policy']['mode']}ed.")
        return result

    if action in ("SCREENSHOT", "CAPTURE"):
        shot_dir = ctx["shots_dir"]
        shot_dir.mkdir(parents=True, exist_ok=True)
        path = shot_dir / f"case_{ctx['index']}_evidence_{int(time.time() * 1000)}.png"
        await page.screenshot(path=str(path))
        result["screenshot"] = str(path)
        return result

    # ── Interactions ──────────────────────────────────────────────────
    if action == "TYPE":
        # Never into a password field: these bookkeeping values are echoed into
        # report.json, which is committed — that echo is how a scrubbed
        # recording still leaked the password into git. A password is also
        # never the created record's identity, so nothing is lost.
        if value.strip() and credential_store.classify_field(step) != "password":
            ctx["runtime_pool"]["last_typed_value"] = value.strip()
            # The first value typed into a form is almost always the record's
            # name/code — the identity any later assertion should verify.
            ctx["runtime_pool"].setdefault("first_typed_value", value.strip())
        cands = candidate_selectors(step, "TYPE")
        last: Exception | None = None
        for attempt in range(3):
            loc, used = await first_usable(cands)
            if loc:
                try:
                    await highlight(loc)
                    await loc.fill(value, timeout=5000 if headless else 8000)
                    if value and await is_keystroke_widget(loc):
                        # fill() pastes the value in one event. Masked number
                        # widgets (PrimeNG p-inputnumber and similar) act only on
                        # keystrokes, so a pasted value skips their formatting,
                        # their letter filter and any calculation bound to them
                        # (Liberty's CBM stayed empty). Type it as a person would.
                        await loc.fill("")
                        await loc.press_sequentially(value, delay=15)
                    result["selector_used"] = used
                    await note_input_rejection(page, loc, step, value, evidence)
                    return result
                except Exception as err:
                    last = err
            if attempt < 2:
                await emit_log(
                    f"🔄 [Retry {attempt + 1}/2] TYPE on '{selector}' failed — retrying..."
                )
                backoff = (ctx.get("pace") or pacing.pace_for("")).retry_backoff_ms
                await page.wait_for_timeout(backoff * (attempt + 1))
        healed = await heuristics.attempt_auto_healing(
            page, "TYPE", selector, value, allow_vision=not is_destructive(step)
        )
        if healed:
            evidence.healed_steps += 1
            result["healed"] = True
            await emit_log(f"🛠️ [Auto-Heal] Repaired input on '{selector}'")
            return result
        # Say why: "not found" and "found but would not accept input" (disabled,
        # read-only, covered) are different defects, and the report is where a
        # tester decides which one this is.
        if last is None:
            raise Exception(f"TYPE failed: '{selector}' was not found on the page.")
        reason = str(last).strip().splitlines()[0][:200]
        raise Exception(f"TYPE failed on '{selector}': the field was found but did "
                        f"not accept input ({reason})")

    if action == "CLEAR":
        loc, used = await first_usable(candidate_selectors(step, "TYPE"))
        if not loc:
            raise Exception(f"CLEAR failed: '{selector}' not found.")
        await loc.fill("")
        result["selector_used"] = used
        return result

    if action in ("CLICK", "DOUBLE_CLICK", "RIGHT_CLICK"):
        cands = candidate_selectors(step, "CLICK")
        for attempt in range(3):
            loc, used = await first_usable(cands)
            if loc:
                try:
                    await highlight(loc)
                    if action == "DOUBLE_CLICK":
                        await loc.dblclick(timeout=5000 if headless else 8000)
                    elif action == "RIGHT_CLICK":
                        await loc.click(button="right", timeout=5000 if headless else 8000)
                    else:
                        await loc.click(timeout=5000 if headless else 8000)
                    result["selector_used"] = used
                    return result
                except Exception:
                    pass
            if attempt < 2:
                await emit_log(
                    f"🔄 [Retry {attempt + 1}/2] CLICK on '{selector}' failed — retrying..."
                )
                backoff = (ctx.get("pace") or pacing.pace_for("")).retry_backoff_ms
                await page.wait_for_timeout(backoff * (attempt + 1))
        healed = await heuristics.attempt_auto_healing(
            page, "CLICK", selector, "", allow_vision=not is_destructive(step)
        )
        if healed:
            evidence.healed_steps += 1
            result["healed"] = True
            await emit_log(f"🛠️ [Auto-Heal] Recovered click on '{selector}'")
            return result
        # ── Self-healing: an application that gates submission ────────
        # Rather than reporting "CLICK failed on Save" and waiting for a human to
        # re-record, work out why the control is disabled and fix what can be
        # fixed safely: fill the required fields this scenario does not own, then
        # try once more.
        if ctx.get("auto_complete", True) and looks_like_submit(step, selector) \
                and not ctx.get("_autocompleted"):
            gated = await is_disabled_now(page, selector)
            if gated:
                ctx["_autocompleted"] = True          # only ever attempt this once
                learned.remember_submit_gating(
                    ctx["project_name"], ctx["flow_name"], "disabled")
                outcome = await autocomplete_form(page, step, ctx["scenario"], ctx, selector)
                evidence.auto_filled.extend(outcome["filled"])

                if outcome["filled"] and not await is_disabled_now(page, selector):
                    loc, used = await first_usable(candidate_selectors(step, "CLICK"))
                    if loc:
                        await highlight(loc)
                        await loc.click(timeout=6000)
                        result["selector_used"] = used
                        result["healed"] = True
                        result["auto_filled"] = outcome["filled"]
                        if used and used != selector:
                            learned.remember_healed(ctx["project_name"], ctx["flow_name"],
                                                     selector, used)
                        evidence.healed_steps += 1
                        await emit_log(
                            f"   ✅ recovered: filled {len(outcome['filled'])} required field(s) "
                            f"the recording had missed, then clicked '{selector}'"
                        )
                        return result

                if outcome["blocked"]:
                    # The scenario deliberately left a mandatory field empty and the
                    # application refused to submit. That refusal IS the validation.
                    evidence.submission_blocked = {
                        "control": selector,
                        "empty_required": outcome["blocked"],
                    }
                    await emit_log(
                        f"   🚧 '{selector}' stayed disabled because "
                        f"{', '.join(str(b) for b in outcome['blocked'])} "
                        f"{'is' if len(outcome['blocked']) == 1 else 'are'} empty — "
                        f"the application blocked the invalid submission."
                    )
                    return result

        raise Exception(await explain_click_failure(page, step, selector))

    if action in ("SELECT", "SELECT_OPTION", "CHOOSE"):
        cands = candidate_selectors(step, "SELECT")
        loc, used = await first_usable(cands)
        # Native <select>
        if loc:
            try:
                tag = (await loc.evaluate("el => el.tagName") or "").lower()
            except Exception:
                tag = ""
            if tag == "select":
                try:
                    await loc.select_option(label=value, timeout=4000)
                except Exception:
                    try:
                        await loc.select_option(value=value, timeout=4000)
                    except Exception:
                        await loc.select_option(index=0, timeout=4000)
                result["selector_used"] = used
                result["selected"] = value
                return result
        # Framework dropdown: open the trigger, then pick the option by text.
        trigger, used = await first_usable(
            [*cands, f"{selector} .p-dropdown-trigger", ".p-dropdown", ".mat-mdc-select",
             ".ant-select", "[role='combobox']"]
        )
        if trigger:
            await highlight(trigger)
            await trigger.click(timeout=4000)
            await page.wait_for_timeout(400)
        option_selectors = [
            f".p-dropdown-item:has-text('{value}')", f"li[role='option']:has-text('{value}')",
            f".mat-mdc-option:has-text('{value}')", f".ant-select-item:has-text('{value}')",
            f"option:has-text('{value}')", f"[role='option']:has-text('{value}')",
            f"li:has-text('{value}')",
        ] if value else [".p-dropdown-item", "li[role='option']", "[role='option']"]
        option, opt_used = await first_usable(option_selectors)
        if not option:
            raise Exception(
                f"SELECT failed: option {value!r} not found in dropdown '{selector}'."
            )
        await option.click(timeout=4000)
        result["selector_used"] = opt_used
        result["selected"] = value
        if value:
            ctx["runtime_pool"]["last_selected_value"] = value
        return result

    if action in ("CHECK", "UNCHECK"):
        loc, used = await first_usable(candidate_selectors(step, "CLICK"))
        if not loc:
            raise Exception(f"{action} failed: '{selector}' not found.")
        try:
            if action == "CHECK":
                await loc.check(timeout=4000)
            else:
                await loc.uncheck(timeout=4000)
        except Exception:
            await loc.click(timeout=4000)
        result["selector_used"] = used
        return result

    if action in ("KEYPRESS", "PRESS", "KEY"):
        key = (value or "Enter").strip() or "Enter"
        key = {"enter": "Enter", "tab": "Tab", "escape": "Escape", "esc": "Escape",
               "space": "Space", "backspace": "Backspace", "delete": "Delete"}.get(
                   key.lower(), key)
        if selector:
            loc, used = await first_usable(candidate_selectors(step, "CLICK"))
            if loc:
                await loc.press(key, timeout=4000)
                result["selector_used"] = used
                result["key"] = key
                return result
        await page.keyboard.press(key)
        result["key"] = key
        return result

    if action in ("UPLOAD", "ATTACH"):
        file_path = value
        if not os.path.isabs(file_path):
            candidate = ctx["flow_path"] / "data" / file_path
            file_path = str(candidate if candidate.exists() else config.BASE_DIR / file_path)
        if not os.path.exists(file_path):
            raise Exception(
                f"UPLOAD failed: file '{value}' not found. Place it in the flow's data/ folder."
            )
        loc, used = await first_usable(
            [*candidate_selectors(step, "UPLOAD"), "input[type='file']"], require_visible=False
        )
        if not loc:
            raise Exception(f"UPLOAD failed: no file input matched '{selector}'.")
        await loc.set_input_files(file_path)
        result["selector_used"] = used
        result["uploaded"] = file_path
        return result

    if action == "HOVER":
        loc, used = await first_usable(candidate_selectors(step, "CLICK"))
        if not loc:
            raise Exception(f"HOVER failed: '{selector}' not found.")
        await loc.hover(timeout=4000)
        result["selector_used"] = used
        return result

    if action in ("SCROLL_TO", "SCROLL"):
        loc, used = await first_usable(candidate_selectors(step, "CLICK"), require_visible=False)
        if loc:
            await loc.scroll_into_view_if_needed(timeout=4000)
            result["selector_used"] = used
        else:
            await page.mouse.wheel(0, 600)
        return result

    # ── STORE: capture a value so later modules can use it ────────────
    # This is what makes a dependency chain real: module 7 generates a code,
    # STORE puts it in the pool, module 8 reads it as {{m7_code}}.
    if action in ("STORE", "CAPTURE_VALUE", "REMEMBER"):
        name = (value or "captured_value").strip().strip("{}")
        mode = (step.get("capture") or "value").lower()
        captured = ""
        if mode == "typed" and ctx["runtime_pool"].get("last_typed_value"):
            captured = str(ctx["runtime_pool"]["last_typed_value"])
        if not captured and selector:
            loc, used = await first_usable(candidate_selectors(step, "TYPE"),
                                           require_visible=False)
            if not loc:
                loc, used = await first_usable(candidate_selectors(step, "CLICK"),
                                               require_visible=False)
            if loc:
                captured = await _read_value(loc)
                result["selector_used"] = used
        if not captured:
            captured = str(ctx["runtime_pool"].get("last_typed_value", "")).strip()
        if not captured:
            raise Exception(
                f"STORE failed: nothing to capture from '{selector}'. Point it at the field "
                f"holding the generated value, or type the value earlier in the scenario."
            )
        ctx["runtime_pool"][name] = captured
        evidence.notes.append(f"Captured {{{{{name}}}}} = {captured!r}")
        await emit_log(f"   ⚓ stored {{{{{name}}}}} = {captured!r}")
        result["stored"] = {name: captured}
        return result

    # ── WAIT_FOR_STATE: poll instead of sleeping ──────────────────────
    # An approval queue moves from Pending to Approved on someone else's clock;
    # a fixed sleep is either flaky or slow, so poll with a deadline.
    if action in ("WAIT_FOR_STATE", "POLL_UNTIL", "WAIT_UNTIL"):
        expected = (value or "").strip()
        timeout_ms = int(step.get("timeout") or step.get("timeout_ms") or 60000)
        interval_ms = int(step.get("interval") or 2000)
        mode = (step.get("match") or "contains").lower()
        deadline = time.time() + timeout_ms / 1000
        seen = ""
        attempts = 0
        while time.time() < deadline:
            attempts += 1
            state: RunState | None = ctx.get("state")
            if state and state.cancelled:
                raise RunCancelled()
            loc, used = await first_usable(candidate_selectors(step, "CLICK"),
                                           require_visible=False)
            if loc:
                seen = (await _read_value(loc) or "").strip()
                result["selector_used"] = used
                if not expected:
                    if seen:
                        break
                elif mode == "equals" and seen.lower() == expected.lower():
                    break
                elif mode != "equals" and expected.lower() in seen.lower():
                    break
            await page.wait_for_timeout(interval_ms)
            # Reloading is opt-in: a single-page app updates the DOM in place, and
            # a reload would discard the very state being waited for. Server-rendered
            # status pages can set {"reload": true}.
            if step.get("reload"):
                try:
                    await page.reload(wait_until="domcontentloaded", timeout=20000)
                    await wait_for_settle(page, 3000)
                except Exception:
                    pass
        else:
            raise Exception(
                f"Timed out after {timeout_ms // 1000}s waiting for '{selector}' to reach "
                f"{expected!r} — it was {seen!r} at the last check ({attempts} polls). "
                f"If this state is set by another user or a background job, confirm the "
                f"expected SLA with the business."
            )
        evidence.notes.append(f"'{selector}' reached {seen!r} after {attempts} poll(s)")
        await emit_log(f"   ⏱ '{selector}' reached {seen!r} after {attempts} poll(s)")
        result["polled"] = {"value": seen, "attempts": attempts}
        return result

    # ── API_REQUEST: reuse the browser session for API-level checks ───
    if action in ("API_REQUEST", "API_GET", "API_POST", "API_PUT", "API_DELETE"):
        method = (step.get("method") or action.replace("API_", "") or "GET").upper()
        if method == "REQUEST":
            method = "GET"
        url = heuristics.resolve_dynamic_variables(selector or step.get("url", ""),
                                                   ctx["runtime_pool"])
        if not url:
            raise Exception("API_REQUEST needs a URL or path.")
        if not url.startswith(("http://", "https://")):
            base = (ctx.get("base_url") or "").rstrip("/")
            root = "/".join(base.split("/")[:3]) if base else ""
            url = root + "/" + url.lstrip("/")
        body = step.get("json") or step.get("body")
        if isinstance(body, (dict, list)):
            body = json.loads(heuristics.resolve_dynamic_variables(
                json.dumps(body), ctx["runtime_pool"]))
        options: dict = {"timeout": 30000}
        if body is not None:
            options["data"] = body
        # page.request shares the page's cookies, so an authenticated UI session
        # authenticates the API call too — no separate token plumbing.
        response = await page.request.fetch(url, method=method, **options)
        text = await response.text()
        try:
            parsed = json.loads(text)
        except Exception:
            parsed = None
        ctx["runtime_pool"]["api_status"] = response.status
        ctx["runtime_pool"]["api_body"] = text[:4000]
        ctx["api_last"] = {"status": response.status, "json": parsed, "text": text,
                           "url": url, "method": method}
        for name, path in (step.get("capture") or {}).items():
            if not isinstance(path, str):
                continue
            found = api_seed._extract(parsed, path) if parsed is not None else None
            if found is not None:
                ctx["runtime_pool"][name] = found
                await emit_log(f"   ⚓ api captured {{{{{name}}}}} = {found!r}")
        evidence.notes.append(f"{method} {url} → HTTP {response.status}")
        await emit_log(f"   🔌 {method} {url} → {response.status}")
        result["api"] = {"status": response.status, "url": url}
        return result

    # ── SWITCH_ROLE: creator, approver and scanner are different people ──
    if action in ("SWITCH_ROLE", "LOGIN_AS"):
        role = (value or selector or "").strip()
        if not role:
            raise Exception("SWITCH_ROLE needs a role name.")
        switcher = ctx.get("switch_role")
        if not switcher:
            raise Exception("Role switching is unavailable in this run.")
        await switcher(role)
        ctx["runtime_pool"]["current_role"] = role
        evidence.notes.append(f"Switched to the '{role}' role")
        result["role"] = role
        return result

    # ── Assertions ────────────────────────────────────────────────────
    if action in ASSERTION_ACTIONS:
        await wait_for_settle(page, 4000)
        await scan_messages(page, evidence, ctx.get("message_selectors"),
                            ctx.get("project_name", ""))
        note_created_record(ctx["runtime_pool"], evidence.messages)
        ok, detail = await run_assertion(page, action, step, selector, value, ctx)
        evidence.assertions.append(
            {"action": action, "selector": selector, "value": value, "ok": ok, "detail": detail}
        )
        result["assertion_ok"] = ok
        result["detail"] = detail
        if not ok:
            # A VERIFY_MESSAGE that failed *only because the wording differs* is
            # not evidence the application misbehaved. The literal in the step is
            # a guess — usually a model's — while the message on screen is what
            # the application really says. Failing the scenario here blames a
            # working app for a wording mistake in its own test, which is the most
            # damaging thing this tool can do. Record the mismatch and let the
            # verdict engine judge from the evidence instead.
            if action == "VERIFY_MESSAGE" and evidence.messages:
                evidence.wording_mismatches.append({
                    "expected": str(value),
                    "actual": [m["text"] for m in evidence.messages][:3],
                    "selector": selector,
                    "step": ctx.get("step_number", 0),
                })
                await emit_log(
                    f"⚠️ {action}: the application said "
                    f"{evidence.messages[0]['text']!r}, this step expected "
                    f"{str(value)!r} — recorded as a wording mismatch, not a failure."
                )
                return result
            raise Exception(f"Validation failed: {detail}")
        await emit_log(f"✅ {action} passed — {detail}")
        return result

    result["skipped"] = f"unsupported action '{action}'"
    await emit_log(f"⚠️ Step skipped: unsupported action '{action}'")
    return result


FORM_INVENTORY_JS = """(selector) => {
__SHARED_PROBES__
    // Everything the form actually contains — not just what a recording touched.
    // This is how the platform learns about a required field nobody recorded.
    let target = null;
    try { target = document.querySelector(selector); } catch (e) { target = null; }

    // Which form are we looking at? The container holding the control we tried to
    // click, if it has one. Otherwise the first *visible* form-like container —
    // applications routinely leave closed dialogs in the DOM ahead of the real
    // content, and taking the first match blindly made this inventory describe a
    // dialog nobody could see, so no required field was ever found.
    const onScreen = (el) => {
        if (!el) return false;
        const style = getComputedStyle(el);
        if (style.display === 'none' || style.visibility === 'hidden') return false;
        const rect = el.getBoundingClientRect();
        return rect.width > 0 && rect.height > 0;
    };
    const CONTAINERS = 'p-dialog, .p-dialog, form, .modal, dialog';
    let scope = null;
    if (target) {
        const owner = target.closest(CONTAINERS);
        if (onScreen(owner)) scope = owner;
    }
    if (!scope) {
        const marked = document.querySelector('[data-qa-surface]');
        if (onScreen(marked)) scope = marked;
    }
    if (!scope) {
        for (const el of document.querySelectorAll(CONTAINERS)) {
            if (onScreen(el)) { scope = el; break; }
        }
    }
    if (!scope) scope = document.body;

    const cssEscape = (v) => (window.CSS && CSS.escape) ? CSS.escape(v) : v;

    const labelFor = (el) => {
        try {
            if (el.id) {
                const l = document.querySelector('label[for="' + cssEscape(el.id) + '"]');
                if (l && l.innerText.trim()) return l.innerText.trim();
            }
            const w = el.closest('.field, .form-group, .p-field, .form-row, .mb-3, .col, .grid');
            const l2 = w && w.querySelector('label');
            if (l2 && l2.innerText.trim()) return l2.innerText.trim();
        } catch (e) {}
        return el.getAttribute('placeholder') || el.getAttribute('formcontrolname') || '';
    };

    const stableSelector = (el) => {
        for (const attr of ['data-testid', 'data-test', 'data-cy']) {
            const v = el.getAttribute(attr);
            if (v) return '[' + attr + '="' + v + '"]';
        }
        if (el.id && !/^(pn_id_|p-|mat-|cdk-|ng-)/i.test(el.id)) return '#' + cssEscape(el.id);
        const fcn = el.getAttribute('formcontrolname');
        if (fcn) return el.tagName.toLowerCase() + '[formcontrolname="' + fcn + '"]';
        if (el.name) return el.tagName.toLowerCase() + '[name="' + el.name + '"]';
        const ph = el.getAttribute('placeholder');
        if (ph) return el.tagName.toLowerCase() + '[placeholder="' + ph + '"]';
        return '';
    };

    const out = [];
    const seen = new Set();

    // 1. Native controls.
    scope.querySelectorAll('input, textarea, select').forEach((el) => {
        if (el.type === 'hidden') return;
        if (!onScreen(el)) return;          // a closed dialog's fields are not fillable
        const sel = stableSelector(el);
        if (!sel || seen.has(sel)) return;
        seen.add(sel);
        const label = labelFor(el);
        const requiredSource = qaRequiredInfo(el, qaLabelElFor(el), label);
        const entry = {
            selector: sel,
            label: label,
            kind: el.tagName.toLowerCase() === 'select' ? 'select'
                  : (el.getAttribute('type') || 'text').toLowerCase(),
            required: !!requiredSource,
            required_source: requiredSource,
            value: (el.value == null ? '' : String(el.value)).trim(),
            disabled: !!el.disabled || el.hasAttribute('readonly'),
        };
        if (el.getAttribute('min') != null) entry.min = el.getAttribute('min');
        if (el.getAttribute('max') != null) entry.max = el.getAttribute('max');
        if (el.getAttribute('maxlength') != null) entry.maxlength = el.getAttribute('maxlength');
        if (entry.kind === 'select') {
            // Skip the placeholder ("-- select --"): choosing it leaves the form
            // invalid, which is exactly the state we are trying to escape.
            entry.options = [...el.options]
                .filter(o => String(o.value || '').trim() !== '')
                .map(o => (o.innerText || o.value || '').trim())
                .filter(Boolean);
        }
        out.push(entry);
    });

    // 2. Component wrappers, whose own element carries the validity state.
    scope.querySelectorAll('p-dropdown, p-inputnumber, p-calendar, p-multiselect,' +
                           ' p-autocomplete, mat-select, nz-select').forEach((el) => {
        const sel = stableSelector(el);
        if (!sel || seen.has(sel)) return;
        seen.add(sel);
        const cls = String(el.className || '');
        const label = labelFor(el);
        const inner = el.querySelector('input');
        out.push({
            selector: sel,
            label: label,
            kind: el.tagName.toLowerCase().includes('inputnumber') ? 'number' : 'dropdown',
            // ng-invalid on an untouched wrapper is the framework saying
            // "required"; the shared probe adds asterisk, CSS marker and classes.
            required: !!(/\\bng-invalid\\b/.test(cls) ||
                         qaRequiredInfo(inner || el, qaLabelElFor(el), label)),
            required_source: /\\bng-invalid\\b/.test(cls) ? 'ng-invalid-empty'
                : qaRequiredInfo(inner || el, qaLabelElFor(el), label),
            value: inner ? (inner.value || '').trim() : '',
            filled: /\\bp-inputwrapper-filled\\b/.test(cls),
            component: true,
        });
    });

    return out;
}"""


# Shared with the recorder so the two can never disagree about which
# fields are mandatory — they did, and that cost four attempted fixes.
FORM_INVENTORY_JS = FORM_INVENTORY_JS.replace("__SHARED_PROBES__", SHARED_PROBES_JS)

READONLY_SCAN_JS = """() => {
    // Inside the open dialog if there is one, otherwise the page.
    const surface = document.querySelector('[data-qa-surface]') || document;
    const label = (el) => {
        const id = el.getAttribute('id');
        if (id) {
            const lab = document.querySelector(`label[for="${id}"]`);
            if (lab && lab.innerText.trim()) return lab.innerText.trim();
        }
        const wrap = el.closest('.field, .form-group, .p-field, td, div');
        const lab = wrap && wrap.querySelector('label');
        if (lab && lab.innerText.trim()) return lab.innerText.trim();
        return el.getAttribute('placeholder') || el.getAttribute('name')
            || el.getAttribute('formcontrolname') || el.tagName.toLowerCase();
    };
    const out = [];
    surface.querySelectorAll('input, textarea, select').forEach((el) => {
        if (el.type === 'hidden' || el.offsetParent === null) return;
        const cls = (el.className || '') + ' ' + ((el.closest('.p-inputwrapper') || {}).className || '');
        const locked = el.readOnly || el.disabled || el.hasAttribute('readonly')
            || /\\breadonly\\b|\\bp-disabled\\b/.test(cls);
        if (locked) out.push(label(el));
    });
    return out;
}"""


EMPTY_REQUIRED_JS = """(selector) => {
__EMPTY_REQUIRED_PROBES__
    // Find the form or dialog the control belongs to, then report which of its
    // mandatory fields are still empty. Covers plain inputs and the component
    // wrappers (PrimeNG / Material / Ant) whose own element holds the state.
    let target = null;
    try { target = document.querySelector(selector); } catch (e) { target = null; }

    // Same rule as the form inventory: never describe a dialog that is closed.
    const onScreen = (el) => {
        if (!el) return false;
        const style = getComputedStyle(el);
        if (style.display === 'none' || style.visibility === 'hidden') return false;
        const rect = el.getBoundingClientRect();
        return rect.width > 0 && rect.height > 0;
    };
    const CONTAINERS = 'p-dialog, .p-dialog, form, .modal, dialog';
    let scope = null;
    if (target) {
        const owner = target.closest(CONTAINERS);
        if (onScreen(owner)) scope = owner;
    }
    if (!scope) {
        const marked = document.querySelector('[data-qa-surface]');
        if (onScreen(marked)) scope = marked;
    }
    if (!scope) {
        for (const el of document.querySelectorAll(CONTAINERS)) {
            if (onScreen(el)) { scope = el; break; }
        }
    }
    if (!scope) scope = document.body;

    const labelFor = (el) => {
        try {
            if (el.id) {
                const l = document.querySelector('label[for="' + CSS.escape(el.id) + '"]');
                if (l && l.innerText.trim()) return l.innerText.trim();
            }
            const w = el.closest('.field, .form-group, .p-field, .form-row, .mb-3, .col, .grid');
            const l2 = w && w.querySelector('label');
            if (l2 && l2.innerText.trim()) return l2.innerText.trim();
        } catch (e) {}
        return el.getAttribute('placeholder') || el.getAttribute('formcontrolname') ||
               el.id || el.tagName.toLowerCase();
    };

    // This used to be a third private detector (end-anchored asterisk plus the
    // two attributes), so "which mandatory fields are still empty" could
    // contradict the inventory that had just listed them. One probe, everywhere.
    const requiredLabel = (el) =>
        !!qaRequiredInfo(el, qaLabelElFor(el), labelFor(el));

    const empty = [];
    // 1. Plain fields the browser knows about.
    scope.querySelectorAll('input, textarea, select').forEach((el) => {
        if (el.type === 'hidden' || el.disabled) return;
        const value = (el.value == null ? '' : String(el.value)).trim();
        if (!value && requiredLabel(el)) empty.push(labelFor(el));
    });
    // 2. Angular/PrimeNG components: ng-invalid on an untouched control is the
    //    framework itself saying "this is required and not filled".
    scope.querySelectorAll('[formcontrolname], p-dropdown, p-inputnumber, p-calendar,' +
                           ' p-multiselect, p-autocomplete').forEach((el) => {
        const cls = String(el.className || '');
        if (!/\\bng-invalid\\b/.test(cls)) return;
        if (/\\bp-inputwrapper-filled\\b/.test(cls)) return;
        // ng-invalid on a FILLED control means "bad value", not "still empty" —
        // the filled-class check above only exists on PrimeNG wrappers, so plain
        // inputs (e.g. a formcontrolname text field holding a rejected value)
        // were reported as empty-mandatory when they were neither.
        const inner = el.matches('input, textarea, select') ? el
                    : el.querySelector('input, textarea, select');
        if (inner && String(inner.value == null ? '' : inner.value).trim()) return;
        const name = labelFor(el);
        if (name && empty.indexOf(name) === -1) empty.push(name);
    });
    return [...new Set(empty)].slice(0, 8);
}"""

# Same shared detector as the inventory and the recorder — this probe used to
# carry a third private copy, so "which mandatory fields are still empty" could
# disagree with the inventory that had just listed them.
EMPTY_REQUIRED_JS = EMPTY_REQUIRED_JS.replace("__EMPTY_REQUIRED_PROBES__", SHARED_PROBES_JS)


# The overlay containers every major component library uses for modals.
SURFACE_MARK = "data-qa-surface"
ACTIVE_SURFACE_JS = """(mark) => {
    const containers = [
        'p-dialog', '.p-dialog', '[role="dialog"]', '[role="alertdialog"]',
        'dialog[open]', '.modal.show', '.modal.in', '.mat-dialog-container',
        '.mat-mdc-dialog-container', '.ant-modal', '.MuiDialog-container',
        '.cdk-overlay-pane', '.p-sidebar', '.offcanvas.show', '.swal2-popup',
    ];
    document.querySelectorAll('[' + mark + ']').forEach(el => el.removeAttribute(mark));

    let best = null;
    let bestScore = -1;
    containers.forEach((sel) => {
        let found = [];
        try { found = document.querySelectorAll(sel); } catch (e) { return; }
        found.forEach((el) => {
            const rect = el.getBoundingClientRect();
            const style = getComputedStyle(el);
            // A real, open dialog: on screen, big enough to hold a form, visible.
            if (rect.width < 80 || rect.height < 60) return;
            if (style.display === 'none' || style.visibility === 'hidden') return;
            if (parseFloat(style.opacity || '1') < 0.1) return;
            if (!el.querySelector('input, select, textarea, button')) return;
            // Topmost wins: z-index first, then depth in the document.
            let z = 0;
            for (let node = el; node; node = node.parentElement) {
                const value = parseInt(getComputedStyle(node).zIndex, 10);
                if (!Number.isNaN(value)) { z = Math.max(z, value); }
            }
            let depth = 0;
            for (let node = el; node; node = node.parentElement) depth++;
            const score = z * 1000 + depth;
            if (score >= bestScore) { bestScore = score; best = el; }
        });
    });

    if (!best) return '';
    best.setAttribute(mark, '1');
    return best.tagName.toLowerCase();
}"""


async def mark_active_surface(page) -> str:
    """Tag the topmost open dialog so selectors can be scoped to it.

    Returns the dialog's tag name, or "" when the page has no modal open.
    Marking is used instead of returning a selector because several dialogs can
    share the same class and only one of them is actually on screen.
    """
    try:
        return await page.evaluate(ACTIVE_SURFACE_JS, SURFACE_MARK)
    except Exception:
        return ""


def scope_to_surface(selector: str) -> str:
    """Rewrite a selector so it only matches inside the marked dialog."""
    sel = (selector or "").strip()
    if not sel or ">>" in sel or sel.startswith(("xpath=", "text=", "/", "(")):
        return ""                      # engine-specific syntax: leave it alone
    if SURFACE_MARK in sel:
        return sel
    # A comma-separated list needs every branch scoped, not just the first.
    parts = [p.strip() for p in sel.split(",") if p.strip()]
    return ", ".join(f"[{SURFACE_MARK}] {part}" for part in parts)


async def is_disabled_now(page, selector: str) -> bool:
    """Is the control present but not clickable? Covers native and CSS-only disabling."""
    try:
        loc = page.locator(selector).first
        if await loc.count() == 0:
            return False
        if await loc.is_disabled():
            return True
        classes = (await loc.get_attribute("class")) or ""
        return "p-disabled" in classes or "disabled" in classes.split()
    except Exception:
        return False


SUBMIT_WORDS = ("save", "submit", "create", "add", "update", "confirm", "ok",
                "apply", "send", "post", "register", "next", "finish")


def looks_like_submit(step: dict, selector: str) -> bool:
    blob = f"{selector} {step.get('value','')} {step.get('label','')}".lower()
    return any(word in blob for word in SUBMIT_WORDS)


def scenario_owns_field(scenario: dict, field: dict) -> bool:
    """True when the scenario itself is exercising this field.

    The line the platform must never cross: a scenario testing "empty name is
    rejected" must not have its name auto-filled, or the test would silently
    prove the opposite of what it claims.
    """
    blob = json.dumps(scenario.get("steps") or []).lower()
    expectation = str(scenario.get("expected_result", "")).lower()
    name = str(scenario.get("scenario_name", "")).lower()
    for token in (field.get("selector"), field.get("label")):
        text = str(token or "").strip().rstrip("*").strip().lower()
        if len(text) < 3:
            continue
        if text in blob:
            return True
        # "Machine Name" in the title/expectation means the scenario is about it.
        words = [w for w in re.split(r"[^a-z0-9]+", text) if len(w) > 3]
        if words and all(w in expectation or w in name for w in words):
            return True
    return False


async def capture_form_inventory(page, ctx: dict, selector: str = "") -> list[dict]:
    """Record the fields on the visible form, whether or not anything went wrong.

    The inventory used to run in exactly one place: after a Save click had
    already failed *and* the button was found disabled. A form that saves
    correctly therefore never produced one — which is why ``learned.json`` stayed
    empty, why auto-complete had nothing to work from, and why fixing detection
    alone never changed what anyone saw. Both halves had to move.

    Pure observation: appends to ``ctx["observed_fields"]`` and touches no
    verdict, evidence or step outcome. A failure is logged and swallowed, because
    learning must never break a run — but it is *logged*, not silent.
    """
    try:
        inventory = await page.evaluate(FORM_INVENTORY_JS, selector or "")
    except Exception as err:      # noqa: BLE001 - observation must not fail a step
        log.info("Form inventory unavailable for %r: %s", selector, err)
        return []
    if not isinstance(inventory, list) or not inventory:
        return []

    seen: set = ctx.setdefault("_field_seen", set())
    fresh: list[dict] = []
    for entry in inventory:
        target = entry.get("selector")
        if not target:
            continue
        # Keyed on required-ness too: a control is only provably ng-invalid while
        # still empty, so the same field can legitimately report differently
        # later. remember_fields ORs the flag, so an early True is never lost.
        key = (target, bool(entry.get("required")))
        if key in seen:
            continue
        seen.add(key)
        fresh.append(entry)

    if fresh:
        ctx.setdefault("observed_fields", []).extend(fresh)
    return fresh


async def autocomplete_form(page, step: dict, scenario: dict, ctx: dict,
                            selector: str) -> dict:
    """Fill required fields the scenario never mentions, so a submit can proceed.

    This is the platform handling an application's own rules instead of asking a
    human to re-record: many forms keep Save disabled until every mandatory field
    is valid, and a generated scenario cannot know about a field the recording
    never captured.

    Returns ``{"filled": [...], "blocked": [...]}`` — ``blocked`` lists required
    fields the scenario deliberately owns, which are left untouched because
    filling them would invalidate the test.
    """
    result: dict = {"filled": [], "blocked": []}
    try:
        inventory = await page.evaluate(FORM_INVENTORY_JS, selector)
    except Exception:
        return result
    if not isinstance(inventory, list):
        return result

    # Everything observed goes into the profile, whether we fill it or not.
    ctx.setdefault("observed_fields", []).extend(inventory)

    token = str(ctx["runtime_pool"].get("run_token", ""))[:6]
    emit_log = ctx["log"]

    for field in inventory:  # noqa: F402 - dataclasses.field is not used here
        if not field.get("required"):
            continue
        if field.get("value") or field.get("filled"):
            continue
        if field.get("disabled"):
            continue
        if learned.is_sensitive(field):
            result["blocked"].append(field.get("label") or field.get("selector"))
            continue
        if scenario_owns_field(scenario, field):
            # The scenario is testing this field — leaving it empty is the point.
            result["blocked"].append(field.get("label") or field.get("selector"))
            continue

        target = field.get("selector")
        if not target:
            continue
        value = learned.synthetic_value(field, token)
        try:
            if field.get("kind") in ("dropdown", "select"):
                await execute_step(page, {"action": "SELECT", "selector": target,
                                          "value": str(value)}, ctx)
            else:
                await execute_step(page, {"action": "TYPE", "selector": target,
                                          "value": str(value)}, ctx)
            result["filled"].append({"selector": target,
                                     "label": field.get("label") or target,
                                     "value": str(value)})
            await emit_log(
                f"   🩹 auto-filled required field {field.get('label') or target!r} "
                f"= {value!r} (the recording never captured it)"
            )
        except Exception as err:
            log.info("Auto-fill failed for %s: %s", target, err)

    if result["filled"]:
        await wait_for_settle(page, 2500)
    return result


async def explain_click_failure(page, step: dict, selector: str) -> str:
    """Turn "CLICK failed" into something a tester can act on."""
    base = f"CLICK failed on '{selector}'"
    try:
        loc = page.locator(selector).first
        if await loc.count() == 0:
            return f"{base} — no element on the page matches that selector."
        disabled = False
        try:
            disabled = await loc.is_disabled()
        except Exception:
            pass
        if not disabled:
            classes = (await loc.get_attribute("class")) or ""
            disabled = "p-disabled" in classes or "disabled" in classes.split()
        if not disabled:
            if not await loc.is_visible():
                return (f"{base} — the element exists but is not visible. It may be behind an "
                        f"overlay, or the dialog it belongs to is not open yet.")
            return base
        empty = await page.evaluate(EMPTY_REQUIRED_JS, selector)
        if empty:
            return (f"{base} — the control is DISABLED because these required fields are "
                    f"still empty: {', '.join(empty)}. Fill them earlier in the scenario, "
                    f"or assert that the button is disabled instead of clicking it.")
        return (f"{base} — the control is DISABLED. The form is probably incomplete or "
                f"unchanged; the application will not enable it until it validates.")
    except Exception:
        return base


COLUMN_QUALITY_JS = """({ table, columns }) => {
    let host = null;
    try { host = document.querySelector(table); } catch (e) { return { error: 'bad selector' }; }
    if (!host) {
        // Fall back to the table inside the dialog that is currently open.
        const dlg = document.querySelector('[data-qa-surface], p-dialog, [role="dialog"], .modal.show');
        host = dlg ? dlg.querySelector('table') : document.querySelector('table');
    }
    if (!host) return { error: 'no table found on the page' };
    if (host.tagName !== 'TABLE') host = host.querySelector('table') || host;

    const headerCells = [...host.querySelectorAll('thead th, thead td')]
        .map(th => (th.innerText || '').trim());
    const bodyRows = [...host.querySelectorAll('tbody tr')]
        .filter(tr => tr.querySelectorAll('td').length > 0 &&
                      !/no records|no data|empty/i.test(tr.innerText || ''));
    if (!bodyRows.length) return { rows: 0, headers: headerCells };

    const norm = (s) => String(s || '').toLowerCase().replace(/[^a-z0-9]/g, '');
    // Default: every column that has a header. Explicit list wins when given.
    const targets = (columns && columns.length)
        ? columns
        : headerCells.filter(Boolean);

    const checked = [];
    const emptyColumns = [];
    targets.forEach((name) => {
        const index = headerCells.findIndex(h => norm(h) === norm(name) ||
                                                 norm(h).includes(norm(name)));
        if (index === -1) return;                       // column not on screen
        checked.push(headerCells[index]);
        let blank = 0;
        bodyRows.forEach((tr) => {
            const cell = tr.querySelectorAll('td')[index];
            if (!cell) { blank++; return; }
            const text = (cell.innerText || '').replace(/[\\s\\u2014\\u2013-]/g, '').trim();
            if (!text && !cell.querySelector('input, img, svg, i, .p-checkbox')) blank++;
        });
        if (blank === bodyRows.length) {
            emptyColumns.push({ column: headerCells[index], blank: blank });
        }
    });
    return { rows: bodyRows.length, headers: headerCells,
             checked: checked, empty_columns: emptyColumns };
}"""


NOOP_ROWS_JS = """({ table, columns }) => {
    let host = null;
    try { host = document.querySelector(table); } catch (e) { return { error: 'bad selector' }; }
    if (!host) {
        const dlg = document.querySelector('[data-qa-surface], p-dialog, [role="dialog"], .modal.show');
        host = dlg ? dlg.querySelector('table') : document.querySelector('table');
    }
    if (!host) return { error: 'no table found on the page' };
    if (host.tagName !== 'TABLE') host = host.querySelector('table') || host;

    const headers = [...host.querySelectorAll('thead th, thead td')]
        .map(th => (th.innerText || '').trim());
    const norm = (s) => String(s || '').toLowerCase().replace(/[^a-z0-9]/g, '');
    const indexOf = (name) => headers.findIndex(h => norm(h) === norm(name) ||
                                                     norm(h).includes(norm(name)));
    const oldIdx = indexOf(columns[0]);
    const newIdx = indexOf(columns[1]);
    if (oldIdx === -1 || newIdx === -1) {
        return { error: 'columns ' + columns.join(' / ') + ' are not in this table' };
    }

    // "2344.0000" and "2344" are the same number; "—" and "" are both nothing.
    const canonical = (raw) => {
        const text = String(raw || '').replace(/[\\u2014\\u2013]/g, '').trim();
        if (!text) return '';
        const asNumber = Number(text.replace(/,/g, ''));
        if (!Number.isNaN(asNumber) && text !== '') return String(asNumber);
        return text.toLowerCase();
    };

    const rows = [...host.querySelectorAll('tbody tr')]
        .filter(tr => tr.querySelectorAll('td').length > Math.max(oldIdx, newIdx));
    const noop = [];
    rows.forEach((tr) => {
        const cells = tr.querySelectorAll('td');
        const before = (cells[oldIdx].innerText || '').trim();
        const after = (cells[newIdx].innerText || '').trim();
        if (canonical(before) === canonical(after)) {
            noop.push({ label: (cells[0].innerText || '').trim().slice(0, 40),
                        old: before, new: after });
        }
    });
    return { rows: rows.length, noop: noop };
}"""


async def read_options(page, step: dict, selector: str) -> list[str]:
    """Every option a dropdown currently offers.

    Native ``<select>`` elements can be read directly. Framework widgets
    (PrimeNG, Material, Ant) render their list only once opened, so the control
    is clicked, the list is read, then it is closed again with Escape so the
    page is left exactly as it was found.
    """
    options: list[str] = []

    # 1. Native select — no interaction needed.
    for cand in candidate_selectors(step, "CLICK"):
        try:
            loc = page.locator(cand).first
            if await loc.count() == 0:
                continue
            tag = (await loc.evaluate("el => el.tagName") or "").upper()
            if tag == "SELECT":
                texts = await loc.evaluate(
                    "el => Array.from(el.options)"
                    ".map(o => (o.textContent || o.value || '').trim())")
                options = [t for t in (texts or []) if t]
                if options:
                    return options
        except Exception:
            continue

    # 2. Framework widget — open it, read the rendered list, close it.
    opened = False
    for cand in candidate_selectors(step, "CLICK"):
        try:
            loc = page.locator(cand).first
            if await loc.count() == 0 or not await loc.is_visible():
                continue
            await loc.click(timeout=3000)
            opened = True
            break
        except Exception:
            continue

    if opened:
        await page.wait_for_timeout(350)
        for option_sel in OPTION_SELECTORS:
            try:
                if option_sel == "option":
                    found = await page.eval_on_selector_all(
                        option_sel,
                        "els => els.map(e => (e.textContent || e.value || '').trim())")
                else:
                    found = await page.eval_on_selector_all(
                        option_sel,
                        "els => els.filter(e => e.offsetParent !== null)"
                        ".map(e => (e.innerText || e.textContent || '').trim())")
                options.extend([f for f in (found or []) if f])
            except Exception:
                continue
        try:
            await page.keyboard.press("Escape")
        except Exception:
            pass

    # De-duplicate while keeping order.
    seen: set[str] = set()
    unique = []
    for opt in options:
        if opt not in seen:
            seen.add(opt)
            unique.append(opt)
    return unique


async def run_assertion(page, action: str, step: dict, selector: str, value: str,
                        ctx: dict) -> tuple[bool, str]:
    """Evaluate one VERIFY_* step. Returns (ok, human-readable detail)."""
    evidence: verdict_engine.Evidence = ctx["evidence"]
    pool = ctx["runtime_pool"]

    async def locator(*, prefer_visible: bool = True):
        """Resolve the assertion's target, preferring an element on screen.

        ``.first`` alone reported "p-dialog exists but is not visible" on any
        application that leaves closed dialogs in the DOM — the assertion was
        looking at the wrong one of several matches.
        """
        fallback = None
        for cand in candidate_selectors(step, "CLICK"):
            try:
                root = page.locator(cand)
                total = await root.count()
            except Exception:
                continue
            if total == 0:
                continue
            if not prefer_visible:
                return root.first
            for index in range(min(total, 12)):
                match = root.nth(index)
                try:
                    if await match.is_visible():
                        return match
                except Exception:
                    continue
            if fallback is None:
                fallback = root.first          # present but hidden — still reportable
        return fallback

    # ── Grid data quality: empty columns and no-op rows ───────────────
    # An audit dialog that lists field changes but leaves DATE/TIME and USER
    # blank is worthless for traceability, yet every step "succeeds" and every
    # row is present, so nothing in a conventional suite notices.
    if action == "VERIFY_COLUMN_NOT_EMPTY":
        wanted = [c.strip() for c in re.split(r"[,;]", _strip_quotes(value)) if c.strip()]
        table_sel = selector or "table"
        try:
            report = await page.evaluate(COLUMN_QUALITY_JS,
                                         {"table": table_sel, "columns": wanted})
        except Exception as err:
            return False, f"Could not inspect '{table_sel}': {err}"
        if report.get("error"):
            return False, report["error"]
        if not report.get("rows"):
            return False, (f"'{table_sel}' has no data rows, so its columns cannot be "
                           f"checked. Open the dialog or grid first.")

        empty = report.get("empty_columns") or []
        if empty:
            names = ", ".join(f"{c['column']} (blank in {c['blank']}/{report['rows']} rows)"
                              for c in empty)
            return False, (
                f"Columns that must always carry data are blank: {names}. "
                f"A history/audit view without a timestamp or a user cannot be used "
                f"for traceability, even though every row is present."
            )
        checked = ", ".join(report.get("checked") or wanted) or "all columns"
        return True, f"Every row in '{table_sel}' has data in {checked}."

    if action == "VERIFY_NO_NOOP_ROWS":
        table_sel = selector or "table"
        pair = [c.strip() for c in re.split(r"[,;]", _strip_quotes(value)) if c.strip()]
        if len(pair) < 2:
            pair = ["Old Value", "New Value"]
        try:
            report = await page.evaluate(NOOP_ROWS_JS,
                                         {"table": table_sel, "columns": pair})
        except Exception as err:
            return False, f"Could not inspect '{table_sel}': {err}"
        if report.get("error"):
            return False, report["error"]
        noop = report.get("noop") or []
        if noop:
            sample = "; ".join(f"{r['label']}: {r['old']!r} → {r['new']!r}" for r in noop[:3])
            return False, (
                f"{len(noop)} of {report.get('rows', 0)} history rows record a change that "
                f"did not happen — {pair[0]} and {pair[1]} are equivalent: {sample}. "
                f"These entries are noise (often a formatting difference such as "
                f"2344.0000 vs 2344) and they hide the real changes."
            )
        return True, f"No no-op rows: every entry in '{table_sel}' records a real change."

    # ── Dropdown membership: the inactive-master exclusion rule ───────
    if action in ("VERIFY_OPTION_PRESENT", "VERIFY_OPTION_ABSENT"):
        wanted = _strip_quotes(value).strip()
        if not wanted:
            return False, f"{action} needs the option text to look for."
        options = await read_options(page, step, selector)
        found = any(wanted.lower() in opt.lower() for opt in options)
        shown = ", ".join(repr(o) for o in options[:12]) or "none"
        if action == "VERIFY_OPTION_PRESENT":
            if found:
                return True, f"'{wanted}' is offered by '{selector}'."
            return False, (
                f"'{wanted}' is not offered by '{selector}'. Options found: {shown}. "
                f"Either the master was never created, or it is not active."
            )
        if not found:
            return True, f"'{wanted}' is correctly not offered by '{selector}'."
        return False, (
            f"'{wanted}' is still offered by '{selector}' even though it should be excluded. "
            f"An inactive or deleted master must not remain selectable downstream."
        )

    # ── API response assertions ───────────────────────────────────────
    if action == "VERIFY_API":
        last = ctx.get("api_last")
        if not last:
            return False, "VERIFY_API used before any API_REQUEST step in this scenario."
        expected_status = step.get("status")
        if expected_status is not None and int(last["status"]) != int(expected_status):
            return False, (f"{last['method']} {last['url']} returned HTTP {last['status']}, "
                           f"expected {expected_status}.")
        path = step.get("path") or selector
        if path:
            actual = api_seed._extract(last.get("json"), path)
            wanted = _strip_quotes(value)
            if wanted:
                wanted_resolved = heuristics.resolve_dynamic_variables(wanted, pool)
                if actual is None:
                    return False, f"'{path}' is not present in the API response."
                if str(wanted_resolved).lower() not in str(actual).lower():
                    return False, (f"API field '{path}' is {actual!r}, expected "
                                   f"{wanted_resolved!r}.")
                return True, f"API field '{path}' is {actual!r} as expected."
            if actual is None:
                return False, f"'{path}' is missing from the API response."
            return True, f"API field '{path}' is present ({actual!r})."
        return True, f"{last['method']} {last['url']} returned HTTP {last['status']}."

    if action in ("VERIFY_STATE", "VERIFY_VISIBLE"):
        loc = await locator()
        if loc is None:
            return False, f"element '{selector}' does not exist on the page"
        visible = await loc.is_visible()
        return visible, (
            f"element '{selector}' is visible" if visible
            else f"element '{selector}' exists but is not visible"
        )

    if action == "VERIFY_HIDDEN":
        loc = await locator()
        if loc is None:
            return True, f"element '{selector}' is absent, as expected"
        visible = await loc.is_visible()
        return (not visible), (
            f"element '{selector}' is hidden, as expected" if not visible
            else f"element '{selector}' is still visible"
        )

    if action == "VERIFY_DISABLED":
        loc = await locator()
        if loc is None:
            return False, f"element '{selector}' not found to verify disabled state"
        if await loc.is_disabled():
            return True, f"'{selector}' is disabled"
        classes = (await loc.get_attribute("class")) or ""
        if "disabled" in classes.lower() or "p-disabled" in classes.lower():
            return True, f"'{selector}' is marked disabled via CSS class"
        if (await loc.get_attribute("readonly")) is not None:
            return True, f"'{selector}' is read-only"
        return False, f"'{selector}' should be disabled but is editable"

    if action == "VERIFY_ENABLED":
        loc = await locator()
        if loc is None:
            return False, f"element '{selector}' not found"
        enabled = await loc.is_enabled()
        return enabled, f"'{selector}' is {'enabled' if enabled else 'disabled'}"

    if action == "VERIFY_READONLY":
        loc = await locator()
        if loc is None:
            # A selector like `input[readonly]` is a probe, not a real element: if
            # nothing matches, the honest answer is what the form actually does,
            # not "selector not found". Inspect it and say which fields are
            # locked — or that none are, which is itself the finding.
            locked = await page.evaluate(READONLY_SCAN_JS)
            if locked:
                names = ", ".join(locked[:6])
                return True, (f"read-only fields on this form: {names}"
                              + (f" (+{len(locked) - 6} more)" if len(locked) > 6 else ""))
            return False, ("no field on this form is read-only or disabled — a "
                           "system-generated value that must not be edited is "
                           "editable here")
        if (await loc.get_attribute("readonly")) is not None:
            return True, f"'{selector}' is read-only"
        classes = ((await loc.get_attribute("class")) or "").lower()
        if "readonly" in classes or "disabled" in classes or "p-disabled" in classes:
            return True, f"'{selector}' is read-only via CSS class"
        if await loc.is_disabled():
            return True, f"'{selector}' is disabled"
        return False, f"'{selector}' should be read-only but accepts input"

    if action in ("VERIFY_VALUE", "ASSERT_VALUE"):
        loc = await locator()
        if loc is None:
            return False, f"element '{selector}' not found to verify value"
        actual = await _read_value(loc)
        expected = _strip_quotes(value)
        ok = expected.lower() in actual.lower()
        return ok, (
            f"'{selector}' contains {expected!r}" if ok
            else f"'{selector}' holds {actual!r}, expected {expected!r}"
        )

    if action == "VERIFY_NON_ZERO":
        loc = await locator()
        if loc is None:
            return False, f"element '{selector}' not found to verify calculation"
        actual = await _read_value(loc)
        match = re.search(r"-?\d+(?:\.\d+)?", actual.replace(",", ""))
        if not match:
            return False, f"'{selector}' has no numeric value (found {actual!r})"
        number = float(match.group(0))
        return number > 0, (
            f"'{selector}' calculated {number}" if number > 0
            else f"'{selector}' calculated {number}, expected greater than 0"
        )

    if action == "VERIFY_TEXT":
        target = selector or "body"
        step_for_text = dict(step, selector=target)
        loc = None
        for cand in candidate_selectors(step_for_text, "CLICK"):
            try:
                candidate_loc = page.locator(cand).first
                if await candidate_loc.count() > 0:
                    loc = candidate_loc
                    break
            except Exception:
                continue
        if loc is None:
            return False, f"element '{target}' not found"
        text = (await loc.inner_text() or "").strip()
        expected = _strip_quotes(value)
        if expected.lower() in text.lower():
            return True, f"'{target}' shows {expected!r}"
        # Every framework words its empty grid differently — "No records found",
        # "No data available", "Nothing to show". Asserting one wording and
        # failing on another says nothing about the application, so accept any
        # recognised empty-state phrase and report which one appeared.
        if looks_like_empty_state(expected):
            actual_phrase = find_empty_state(text)
            if actual_phrase:
                return True, (f"the grid reported it is empty — it shows "
                              f"{actual_phrase!r} rather than {expected!r}")
        return False, f"'{target}' does not contain {expected!r}"

    if action in ("VERIFY_ROW_CONTAINS", "VERIFY_TABLE_ROW"):
        expected = _strip_quotes(value).strip()
        if not expected or expected.lower() in (
            "expected record name/code", "expected record name", "record_name",
            "record name", "new record",
        ):
            expected = pool.get("last_typed_value", expected)
        found, detail = await find_row(page, selector, expected)
        if found:
            return True, detail

        # A generated scenario often asserts on a literal copied from the original
        # recording ('asdfdsfds'), which this run never created — so the row is
        # legitimately absent and the failure says nothing about the application.
        # Verify against the record THIS scenario created instead, and say so.
        own = [pool.get("created_record"), pool.get("first_typed_value"),
               pool.get("last_typed_value")]
        for candidate in own:
            candidate = str(candidate or "").strip()
            if not candidate or candidate.lower() == expected.lower():
                continue
            found_own, detail_own = await find_row(page, selector, candidate)
            if found_own:
                return True, (
                    f"{expected!r} is not in the table — that value came from the original "
                    f"recording, not from this run. Verified instead against the record this "
                    f"scenario actually created, {candidate!r}: {detail_own}"
                )
        return False, detail

    if action == "VERIFY_ROW_ABSENT":
        expected = _strip_quotes(value).strip() or pool.get("last_typed_value", "")
        found, detail = await find_row(page, selector, expected, max_pages=3, use_search=True)
        return (not found), (
            f"{expected!r} is no longer listed, as expected" if not found
            else f"{expected!r} is still present in the table ({detail})"
        )

    if action == "VERIFY_URL":
        expected = _strip_quotes(value)
        current = page.url
        ok = expected.lower() in current.lower()
        return ok, (
            f"URL contains {expected!r}" if ok
            else f"URL is {current!r}, expected it to contain {expected!r}"
        )

    if action == "VERIFY_COUNT":
        loc_count = 0
        for cand in candidate_selectors(step, "CLICK"):
            try:
                loc_count = await page.locator(cand).count()
                if loc_count:
                    break
            except Exception:
                continue
        expected_raw = _strip_quotes(value)
        match = re.search(r"\d+", expected_raw or "")
        if match:
            expected_n = int(match.group(0))
            if "at least" in expected_raw.lower() or ">" in expected_raw:
                return loc_count >= expected_n, f"found {loc_count} elements (expected ≥ {expected_n})"
            return loc_count == expected_n, f"found {loc_count} elements (expected {expected_n})"
        return loc_count > 0, f"found {loc_count} matching elements"

    if action == "VERIFY_DOWNLOAD":
        # The download handler saves the file in a background task, so sampling
        # `evidence.downloads` once races it: the assertion intermittently
        # reported "no file was downloaded" for an export that worked. Wait for
        # the file to land instead of guessing.
        downloads = evidence.downloads
        if not downloads:
            for _ in range(20):
                await page.wait_for_timeout(250)
                if evidence.downloads:
                    break
            downloads = evidence.downloads
        if not downloads:
            return False, ("no file was downloaded during this scenario (waited 5s "
                           "after the export click)")
        expected_rows = None
        match = re.search(r"\d+", value or "")
        if match:
            expected_rows = int(match.group(0))
        latest = downloads[-1]
        if expected_rows is not None and latest.get("row_count") is not None:
            ok = latest["row_count"] >= expected_rows
            return ok, (
                f"downloaded {latest.get('suggested_filename')} with "
                f"{latest['row_count']} data rows (expected ≥ {expected_rows})"
            )
        return True, (
            f"downloaded {latest.get('suggested_filename')} "
            f"({latest.get('row_count', 'unknown')} rows)"
        )

    if action == "VERIFY_MESSAGE":
        expected = _strip_quotes(value)
        pool_messages = [m["text"] for m in evidence.messages]
        if not pool_messages:
            return False, "no validation or toast message appeared on screen"
        if expected:
            ok = any(expected.lower() in m.lower() for m in pool_messages)
            return ok, (
                f"message {expected!r} shown" if ok
                else f"messages on screen were {pool_messages[:3]}, expected {expected!r}"
            )
        return True, f"message shown: {pool_messages[0]!r}"

    return False, f"unsupported assertion '{action}'"


async def _read_value(loc) -> str:
    for reader in ("input_value", "inner_text"):
        try:
            text = await getattr(loc, reader)()
            if text and str(text).strip():
                return str(text).strip()
        except Exception:
            continue
    try:
        attr = await loc.get_attribute("value")
        if attr:
            return attr.strip()
    except Exception:
        pass
    return ""


# ──────────────────────────────────────────────────────────────────────
# Download bookkeeping
# ──────────────────────────────────────────────────────────────────────
def count_export_rows(path: Path) -> int | None:
    """Row count for a downloaded CSV/XLSX so export tests can assert totals."""
    suffix = path.suffix.lower()
    try:
        if suffix in (".csv", ".txt"):
            with open(path, "r", encoding="utf-8-sig", errors="replace", newline="") as fh:
                rows = [r for r in csv.reader(fh) if any(c.strip() for c in r)]
            return max(len(rows) - 1, 0)
        if suffix in (".xlsx", ".xlsm"):
            try:
                from openpyxl import load_workbook
            except Exception:
                return None
            workbook = load_workbook(filename=str(path), read_only=True)
            total = 0
            for sheet in workbook.worksheets:
                total += max((sheet.max_row or 1) - 1, 0)
            workbook.close()
            return total
    except Exception:
        return None
    return None


# ──────────────────────────────────────────────────────────────────────
# Prerequisite chains and authentication
# ──────────────────────────────────────────────────────────────────────
AUTH_KEYWORDS = ("login", "auth", "session", "logout", "signin", "sign-in")


def is_auth_flow(flow_name: str) -> bool:
    return any(keyword in (flow_name or "").lower() for keyword in AUTH_KEYWORDS)


async def replay_dependencies(page, cfg: RunConfig, ctx: dict, emit_log,
                              deps: list[str] | None = None) -> None:
    """Replay prerequisite flows (login → master A → master B → target).

    ``deps`` lets the caller replay a subset — the suite-level seeding phase
    handles ``run_once`` masters, leaving only the per-scenario ones here.
    """
    if deps is None:
        deps = contracts.dependency_list(cfg.dependencies)
    if not deps:
        return

    token = ctx["runtime_pool"].setdefault("run_token", unique_token())
    for dep in deps:
        dep_path = config.flow_dir(cfg.project_name, dep) / "actions.json"
        if not dep_path.exists():
            await emit_log(f"⚠️ Prerequisite '{dep}' has no recording — skipped.")
            continue

        await emit_log(f"🔗 Resolving prerequisite dependency: [{dep}]...")
        try:
            with open(dep_path, "r", encoding="utf-8") as fh:
                dep_actions = json.load(fh)
        except Exception as err:
            await emit_log(f"⚠️ Could not read prerequisite '{dep}': {err}")
            continue

        make_unique = cfg.dep_unique_data and not is_auth_flow(dep)
        if make_unique:
            await emit_log(
                f"🎲 Unique-data mode: values typed during [{dep}] get the suffix "
                f"'{token}' so repeat runs don't collide with existing records."
            )

        for step_no, act in enumerate(dep_actions, start=1):
            if ctx["state"].cancelled:
                raise RunCancelled()
            action = (act.get("action") or "").upper()
            selector = act.get("selector")
            value = act.get("value", "")
            if not action or not selector:
                continue
            if make_unique and action == "TYPE":
                new_value = make_dep_value_unique(selector, value, token)
                if new_value != value:
                    ctx["runtime_pool"][f"dep_{dep}_{step_no}"] = new_value
                    ctx["runtime_pool"]["last_unique_value"] = new_value
                    value = new_value

            await emit_log(f"🔗 [{dep}] Step {step_no}: {action} on '{selector}'")
            try:
                await execute_step(
                    page,
                    {"action": action, "selector": selector, "value": value,
                     "selectors": act.get("selectors")},
                    ctx,
                )
            except Exception as err:
                await emit_log(
                    f"⚠️ Prerequisite warning on '{selector}': "
                    f"{humanize_error(err, action, selector)}"
                )
            await wait_for_settle(page, 3000)

        # Capture what this prerequisite produced, so the dependent module can
        # reference it by name instead of hardcoding a guess.
        for capture in contracts.capture_steps(contracts.load(cfg.project_name, dep)):
            try:
                await execute_step(page, capture, ctx)
            except Exception as err:
                await emit_log(
                    f"⚠️ Could not capture '{capture.get('value')}' from [{dep}]: "
                    f"{humanize_error(err)}. Check the 'name_field' in its contract."
                )
        await page.wait_for_timeout(1000)


async def seed_prerequisites(browser, cfg: RunConfig, shared: dict, emit: Emit,
                             state: RunState, deps: list[str]) -> dict:
    """Create the run-once master records and return the variables they produced.

    Runs in its own context so a seeding failure can never leave a scenario's
    browser in a strange state, and saves the resulting session so scenarios
    start already authenticated.
    """
    session_file: Path = shared["session_file"]
    context_args: dict[str, Any] = {"accept_downloads": True}
    if session_file.exists():
        context_args["storage_state"] = str(session_file)
    if not cfg.headless:
        context_args["no_viewport"] = True

    context = await browser.new_context(**context_args)
    page = await context.new_page()
    await pacing.track(page)
    page.set_default_timeout(15000)
    await install_message_observer(page)

    evidence = verdict_engine.Evidence()
    pool: dict = {"run_token": unique_token()}
    # A dependency flow's recording is scrubbed the same way the target's is,
    # so its {{username}}/{{password}} placeholders need the pool too.
    pool.update({key: value
                 for key, value in (shared.get("credentials") or {}).items()
                 if key in ("username", "password") and value})

    async def emit_log(message: str) -> None:
        await emit({"event": "log", "message": message})

    ctx = {
        "evidence": evidence,
        "runtime_pool": pool,
        "headless": cfg.headless,
        "pace": pacing.pace_for(cfg.speed),
        "log": emit_log,
        "dialog_policy": {"mode": "accept"},
        "flow_path": shared["flow_path"],
        "shots_dir": shared["flow_path"] / "screenshots",
        "index": 0,
        "base_url": shared["target_url"],
        "state": state,
        "project_name": cfg.project_name,
        "flow_name": cfg.flow_name,
        "scenario": {},
        "auto_complete": cfg.auto_complete,
    }

    try:
        await page.goto(shared["target_url"], wait_until="domcontentloaded", timeout=60000)
        await wait_for_settle(page)
        await replay_authentication(
            page, context, cfg, ctx, emit_log, shared["target_url"],
            shared["credentials"], session_file,
        )
        # Unique data is essential here: the seed runs on every execution, so
        # fixed values would collide with the records the last run created.
        seeding_cfg = replace(cfg, dep_unique_data=True)
        await replay_dependencies(page, seeding_cfg, ctx, emit_log, deps)
        try:
            await context.storage_state(path=str(session_file))
        except Exception:
            pass
    finally:
        try:
            await context.close()
        except Exception:
            pass

    # Publish only the named contract variables, not the internal bookkeeping.
    published = {}
    for dep in deps:
        for provide in contracts.load(cfg.project_name, dep)["provides"]:
            var = provide["var"]
            if pool.get(var):
                published[var] = pool[var]
    published["run_token"] = pool["run_token"]
    return published


def seedable_dependencies(cfg: RunConfig) -> tuple[list[str], list[str]]:
    """Split prerequisites into those seeded once per run and those replayed per scenario.

    A master module marked ``run_once`` creates a record that persists in the
    application, so replaying it for all 30 scenarios only produces 30 duplicates
    and 30× the runtime.
    """
    once: list[str] = []
    each: list[str] = []
    for dep in contracts.dependency_list(cfg.dependencies):
        contract = contracts.load(cfg.project_name, dep)
        if cfg.seed_once and contract.get("run_once") and contract.get("provides") \
                and not is_auth_flow(dep):
            once.append(dep)
        else:
            each.append(dep)
    return once, each


async def replay_authentication(page, context, cfg: RunConfig, ctx: dict, emit_log,
                                target_url: str, credentials: dict,
                                session_file: Path) -> bool:
    """Log in when the target URL bounced us to a login screen."""
    try:
        on_login_page = "auth/login" in page.url or "login" in page.url.lower() or \
            await page.locator("input[name='username'], input[type='password']").count() > 0
    except Exception:
        on_login_page = False

    testing_auth_itself = (cfg.flow_name == cfg.auth_flow) or is_auth_flow(cfg.flow_name) or \
        "auth/login" in (target_url or "")

    if not on_login_page:
        if not testing_auth_itself and not cfg.dependencies.strip():
            await emit_log("🟢 Active login session detected. Bypassing auth flow.")
        return False

    if testing_auth_itself or not cfg.auth_flow or not cfg.auth_scenario:
        if not testing_auth_itself:
            await emit_log(
                "⚠️ Landed on a login page but no authentication profile is configured. "
                "Set 'Authentication Required' in the execution profile, or add 'login' "
                "to Chain Prerequisites."
            )
        return False

    await emit_log(f"🔑 Redirect detected. Executing auth scenario: "
                   f"[{cfg.auth_flow} → {cfg.auth_scenario}]")

    auth_cases_path = config.flow_dir(cfg.project_name, cfg.auth_flow) / "test_cases.json"
    auth_steps: list = []
    if auth_cases_path.exists():
        try:
            with open(auth_cases_path, "r", encoding="utf-8") as fh:
                auth_matrix = json.load(fh)
            auth_case = next(
                (tc for tc in auth_matrix
                 if (tc.get("scenario_name") or tc.get("scenario")) == cfg.auth_scenario),
                None,
            )
            if auth_case:
                auth_steps = auth_case.get("steps", [])
        except Exception as err:
            await emit_log(f"⚠️ Could not read the auth scenario: {err}")

    if not auth_steps:
        actions_path = config.flow_dir(cfg.project_name, cfg.auth_flow) / "actions.json"
        if actions_path.exists():
            try:
                with open(actions_path, "r", encoding="utf-8") as fh:
                    auth_steps = json.load(fh)
                await emit_log("🔑 No matching auth scenario; replaying the raw login recording.")
            except Exception:
                pass

    for step_no, step in enumerate(parse_steps(auth_steps), start=1):
        if ctx["state"].cancelled:
            raise RunCancelled()
        action = (step.get("action") or "").upper()
        selector = step.get("selector") or ""
        value = step.get("value", "")
        if action == "NAVIGATE":
            continue
        if not selector:
            continue
        if action == "TYPE":
            low = selector.lower()
            if "user" in low or "email" in low or "login" in low:
                value = credentials.get("username", value)
            elif "pass" in low or "pwd" in low:
                value = credentials.get("password", value)

        await emit_log(f"🔑 Replaying Auth Step {step_no}: {action} on '{selector}'")
        try:
            await execute_step(page, dict(step, value=value), ctx)
        except Exception as err:
            await emit_log(f"⚠️ Auth step warning: {humanize_error(err, action, selector)}")
        pace = ctx.get("pace") or pacing.pace_for("")
        await wait_for_settle(page, pace.settle_max_ms, pace.quiet_ms)

    await wait_for_settle(page, 6000)

    try:
        still_login = "auth/login" in page.url or \
            await page.locator("input[type='password']").count() > 0
    except Exception:
        still_login = False

    if still_login:
        await emit_log(
            "⛔ Authentication did not complete — still on the login screen. Check the "
            "credentials in the project settings (or login_data.json)."
        )
        return False

    try:
        session_file.parent.mkdir(parents=True, exist_ok=True)
        await context.storage_state(path=str(session_file))
        await emit_log(f"✅ Auth successful! Session cookies saved to {session_file.name}.")
    except Exception as err:
        await emit_log(f"⚠️ Could not persist session cookies: {err}")

    await emit_log(f"Navigating back to target URL: {target_url}")
    await page.goto(target_url, wait_until="domcontentloaded", timeout=60000)
    await wait_for_settle(page)
    return True


# ──────────────────────────────────────────────────────────────────────
# Scenario execution
# ──────────────────────────────────────────────────────────────────────
async def run_scenario(browser, scenario: dict, cfg: RunConfig, shared: dict,
                       emit: Emit, state: RunState) -> dict:
    """Execute one scenario in an isolated browser context and judge the result."""
    index = scenario.get("_original_id", scenario.get("_run_seq", 0))
    label = scenario.get("_label", "")
    name = scenario.get("scenario_name") or scenario.get("scenario") or f"Test Case #{index}"
    expected = scenario.get("expected_result") or scenario.get("expected") or ""
    flow_path: Path = shared["flow_path"]
    target_url: str = shared["target_url"]
    prefix = f"#{index} " if cfg.workers > 1 else ""

    async def emit_log(message: str) -> None:
        await emit({"event": "log", "index": index, "message": f"{prefix}{message}"})

    await emit({
        "event": "test_start",
        "index": index,
        "run_seq": scenario.get("_run_seq", index),
        "total_run": shared["total"],
        "name": name + (f" [{label}]" if label else ""),
        "type": scenario.get("type", "functional"),
    })

    evidence = verdict_engine.Evidence()
    # Message surfaces this application has already been seen to use. Written by
    # remember_message_selector() on every run, so the scanner gets sharper the
    # more the project is exercised — no framework list can cover every app.
    known_message_selectors = learned.message_selectors(cfg.project_name)
    console_errors: list[str] = []
    network_errors: list[str] = []
    steps_executed: list[str] = []
    step_durations: list[dict] = []
    # Credentials first, so a recorded step scrubbed to {{username}} /
    # {{password}} resolves back to the real login. First rather than last:
    # a data-table row supplying its own username (an invalid-login case)
    # must override the project credential, not the other way round.
    runtime_pool: dict = {
        key: value
        for key, value in (shared.get("credentials") or {}).items()
        if key in ("username", "password") and value
    }
    runtime_pool.update(shared.get("seed_pool") or {})
    runtime_pool.update(shared.get("seeded_vars") or {})
    # A scrubbed data row says `username: "{{username}}"`. Adopting that would
    # make the placeholder resolve to itself and get typed literally, so a
    # self-referential value defers to the credential already in the pool.
    for key, value in (scenario.get("_data_row") or {}).items():
        if str(value) == "{{" + str(key) + "}}" and key in runtime_pool:
            continue
        runtime_pool[key] = value
    dialog_policy = {"mode": infer_dialog_policy(scenario)}
    downloads_dir = flow_path / "downloads"
    shots_dir = flow_path / "screenshots"
    start_time = time.time()

    context_args: dict[str, Any] = {"accept_downloads": True}
    session_file: Path = shared["session_file"]
    if session_file.exists():
        context_args["storage_state"] = str(session_file)
    if not cfg.headless:
        context_args["no_viewport"] = True
    video_dir = flow_path / "videos"
    video_dir.mkdir(parents=True, exist_ok=True)
    context_args["record_video_dir"] = str(video_dir)
    # Sized to this screen's page (measured at the start of the run), so the
    # recording fills the frame instead of a corner of it (app/screens.py).
    context_args["record_video_size"] = (shared.get("video_sizes") or {}).get(
        id(browser), dict(screens.DEFAULT_VIDEO))

    context = await browser.new_context(**context_args)
    trace_path = ""
    if cfg.tracing:
        try:
            await context.tracing.start(screenshots=True, snapshots=True, sources=False)
        except Exception:
            pass

    page = await context.new_page()
    await pacing.track(page)
    page.set_default_timeout(15000 if not cfg.headless else 10000)
    await install_message_observer(page)

    page.on("console", lambda msg: console_errors.append(f"[{msg.type}] {msg.text}")
            if msg.type in ("error", "warning") else None)
    page.on("requestfailed", lambda req: network_errors.append(
        f"Network error: {req.method} {req.url} - {req.failure or 'failed'}"))

    def on_dialog(dialog) -> None:
        record = {"type": dialog.type, "message": dialog.message, "action": dialog_policy["mode"]}
        evidence.dialogs.append(record)

        async def resolve() -> None:
            try:
                if dialog_policy["mode"] == "dismiss":
                    await dialog.dismiss()
                else:
                    await dialog.accept()
            except Exception:
                pass
        asyncio.create_task(resolve())

    page.on("dialog", on_dialog)

    def on_download(download) -> None:
        async def save() -> None:
            try:
                downloads_dir.mkdir(parents=True, exist_ok=True)
                suggested = download.suggested_filename or f"download_{int(time.time())}"
                dest = downloads_dir / f"case_{index}_{suggested}"
                await download.save_as(str(dest))
                rows = count_export_rows(dest)
                evidence.downloads.append({
                    "suggested_filename": suggested,
                    "path": str(dest),
                    "url": config.static_url(cfg.project_name, cfg.flow_name,
                                             "downloads", dest.name),
                    "row_count": rows,
                    "size_bytes": dest.stat().st_size if dest.exists() else 0,
                })
            except Exception as err:
                log.warning("Download capture failed: %s", err)
        asyncio.create_task(save())

    page.on("download", on_download)

    ctx = {
        "evidence": evidence,
        "runtime_pool": runtime_pool,
        "headless": cfg.headless,
        "pace": pacing.pace_for(cfg.speed),
        "log": emit_log,
        "dialog_policy": dialog_policy,
        "flow_path": flow_path,
        "shots_dir": shots_dir,
        "index": index,
        "base_url": target_url,
        "state": state,
        "project_name": cfg.project_name,
        "flow_name": cfg.flow_name,
        "scenario": scenario,
        "auto_complete": cfg.auto_complete,
        "message_selectors": known_message_selectors,
    }

    async def switch_role(role: str) -> None:
        """Re-authenticate this context as a different person.

        Module 6 is filled by a creator, module 9 approved by an approver and
        module 10 scanned by an operator — a single credential set cannot test
        that chain honestly.
        """
        creds = resolve_credentials(cfg.project_name, role)
        if not creds.get("username"):
            raise Exception(
                f"No credentials configured for the '{role}' role. Add it under "
                f"Settings → Roles for this project."
            )
        await emit_log(f"👤 Switching to the '{role}' role ({creds['username']})…")
        try:
            await context.clear_cookies()
        except Exception:
            pass
        await page.goto(target_url, wait_until="domcontentloaded", timeout=60000)
        await wait_for_settle(page)
        signed_in = await replay_authentication(
            page, context, cfg, ctx, emit_log, target_url, creds,
            config.project_session_file(cfg.project_name, role),
        )
        if not signed_in:
            raise Exception(f"Could not sign in as '{role}'.")

    ctx["switch_role"] = switch_role

    status = verdict_engine.PASS
    verdict_obj: verdict_engine.Verdict | None = None
    error_text = ""
    screenshot_rel = ""
    temp_video_path = ""
    failed_step_no = 0

    try:
        if state.cancelled:
            raise RunCancelled()

        await emit_log(f"Navigating to target URL: {target_url}")
        await page.goto(target_url, wait_until="domcontentloaded", timeout=60000)
        await wait_for_settle(page)
        evidence.url_start = page.url

        _seeded_once, per_scenario_deps = seedable_dependencies(cfg)
        await replay_dependencies(page, cfg, ctx, emit_log, per_scenario_deps)

        if not cfg.dependencies.strip():
            await replay_authentication(
                page, context, cfg, ctx, emit_log, target_url,
                shared["credentials"], session_file,
            )
        elif session_file.exists():
            try:
                if "auth/login" in page.url:
                    await emit_log(
                        "⚠️ The saved session looks expired (still on the login page)."
                    )
            except Exception:
                pass

        evidence.url_start = page.url
        run_steps = parse_steps(scenario.get("steps", []))
        evidence.steps_total = len(run_steps)

        for step_no, step in enumerate(run_steps, start=1):
            if state.cancelled:
                raise RunCancelled()
            action = (step.get("action") or "").upper()
            selector = step.get("selector") or ""
            value = heuristics.resolve_dynamic_variables(
                str(step.get("value") or ""), runtime_pool
            )
            if not action:
                action = "TYPE" if value else "CLICK"
            if not selector and action not in (
                "KEYPRESS", "PRESS", "KEY", "WAIT", "SLEEP", "DIALOG", "SCREENSHOT",
                "CAPTURE", "VERIFY_URL", "VERIFY_DOWNLOAD", "VERIFY_MESSAGE", "NAVIGATE",
                "SCROLL", "SCROLL_TO",
            ):
                continue

            failed_step_no = step_no
            step_desc = f"{action} on '{selector}'" + (
                f" with value '{value}'" if action in ("TYPE", "SELECT", "KEYPRESS") and value
                else ""
            )
            steps_executed.append(step_desc)
            await emit({"event": "log", "index": index,
                        "message": f"{prefix}Replaying Step {step_no}: {step_desc}"})

            step_start = time.time()
            step_result = await execute_step(page, dict(step, action=action), ctx)
            elapsed_ms = int((time.time() - step_start) * 1000)
            step_durations.append({
                "step": step_no, "action": action, "selector": selector,
                "duration_ms": elapsed_ms, "healed": step_result.get("healed", False),
                "skipped": step_result.get("skipped", ""),
            })
            evidence.steps_executed = step_no

            # A form is on screen most often right after a click or a navigation,
            # so that is when its fields are worth reading — including on the
            # positive flow, where nothing fails and the old trigger (a rejected
            # click on a disabled Save) never fired at all.
            if action in ("CLICK", "NAVIGATE", "SELECT") and not step_result.get("skipped"):
                found = await capture_form_inventory(page, ctx, selector or "")
                mandatory = [f for f in found if f.get("required")]
                if mandatory and not ctx.get("_announced_required"):
                    ctx["_announced_required"] = True
                    await emit({"event": "log", "index": index, "message": (
                        f"{prefix}📋 Form inventory: {len(found)} field(s), "
                        f"{len(mandatory)} mandatory — " + ", ".join(
                            f"{f.get('label') or f.get('selector')} "
                            f"[{f.get('required_source')}]" for f in mandatory[:6])
                        + ("…" if len(mandatory) > 6 else ""))})

            fresh = await scan_messages(page, evidence, known_message_selectors,
                                        cfg.project_name)
            if fresh:
                await emit_log(f"🔍 Screen message: {fresh}")

            if cfg.step_shots == "all":
                try:
                    shots_dir.mkdir(parents=True, exist_ok=True)
                    await page.screenshot(
                        path=str(shots_dir / f"case_{index}_step_{step_no}.png")
                    )
                except Exception:
                    pass

            pace = ctx["pace"]
            if pace.step_pause_ms:
                await page.wait_for_timeout(pace.step_pause_ms if cfg.headless
                                            else pace.step_pause_ms + 300)
            else:
                await wait_for_settle(page, pace.settle_max_ms, pace.quiet_ms)

        await wait_for_settle(page, 5000)
        await scan_messages(page, evidence, known_message_selectors, cfg.project_name)
        evidence.url_end = page.url

        if evidence.messages:
            await emit_log(
                f"🔍 Live Screen Inspector: {[m['text'] for m in evidence.messages][:5]}"
            )

        shots_dir.mkdir(parents=True, exist_ok=True)
        result_shot = shots_dir / f"result_test_{index}.png"
        await page.screenshot(path=str(result_shot))
        screenshot_rel = config.static_url(cfg.project_name, cfg.flow_name,
                                           "screenshots", result_shot.name)

        # ── Verdict ───────────────────────────────────────────────────
        if cfg.verdict_mode == "off":
            verdict_obj = verdict_engine.Verdict(
                verdict_engine.PASS, "All steps executed (verdict checking disabled).",
                "steps", 0.3, evidence.to_dict(),
            )
        else:
            # "Message surface known" means the engine has some grounds for
            # believing it can read this application's messages: it has learned a
            # selector before, or it saw a message during this very scenario.
            surface_known = bool(known_message_selectors) or bool(evidence.messages)
            verdict_obj = verdict_engine.evaluate(expected, evidence,
                                                  scenario.get("type", ""),
                                                  message_surface_known=surface_known,
                                                  declared_intent=scenario.get(
                                                      "expected_intent", ""))
            if cfg.verdict_mode == "ai" and verdict_obj.status == verdict_engine.UNVERIFIED:
                try:
                    shot_bytes = result_shot.read_bytes()
                except Exception:
                    shot_bytes = None
                ai_verdict = await verdict_engine.ai_adjudicate(
                    name, expected, steps_executed, evidence, shot_bytes
                )
                if ai_verdict:
                    verdict_obj = ai_verdict

        # ── Your own rules, checked independently of the verdict ──────
        # A scenario can satisfy its own expectation and still break a rule you
        # stated about the application; that is a failure worth surfacing.
        if cfg.apply_rules:
            rule_results = rules_engine.evaluate(
                shared.get("rules") or [], evidence, scenario, run_steps)
            evidence.rule_results = rule_results
            broken = rules_engine.violations(rule_results)
            for entry in rule_results:
                if entry["skipped"]:
                    continue
                await emit_log(
                    f"   {'📏' if entry['ok'] else '🚫'} Rule “{entry['title']}”: "
                    f"{'holds' if entry['ok'] else 'VIOLATED'} — {entry['detail']}"
                )
            if broken:
                worst = min(broken, key=lambda r: rules_engine.SEVERITIES.index(r["severity"]))
                verdict_obj = verdict_engine.Verdict(
                    verdict_engine.FAIL,
                    f"Rule “{worst['title']}” was violated: {worst['detail']}"
                    + (f" (+{len(broken) - 1} more)" if len(broken) > 1 else ""),
                    "rules", 0.95, evidence.to_dict(),
                )

        status = verdict_obj.status
        await emit_log(
            f"{'✅' if status == verdict_engine.PASS else '⚠️' if status == verdict_engine.UNVERIFIED else '❌'}"
            f" Verdict [{status}] — {verdict_obj.reason}"
        )

    except RunCancelled:
        status = "CANCELLED"
        error_text = "Run stopped by the user."
        verdict_obj = verdict_engine.Verdict(status, error_text, "steps", 1.0, evidence.to_dict())
    except Exception as err:
        raw = str(err)
        is_crash = any(marker in raw.lower() for marker in CRASH_MARKERS)
        status = "CRASH" if is_crash else verdict_engine.FAIL
        error_text = humanize_error(err)
        verdict_obj = verdict_engine.Verdict(
            status,
            f"Step {failed_step_no} failed: {error_text}" if failed_step_no else error_text,
            "steps", 1.0, evidence.to_dict(),
        )
        try:
            shots_dir.mkdir(parents=True, exist_ok=True)
            fail_shot = shots_dir / f"failed_test_{index}.png"
            await page.screenshot(path=str(fail_shot))
            screenshot_rel = config.static_url(cfg.project_name, cfg.flow_name,
                                               "screenshots", fail_shot.name)
        except Exception:
            screenshot_rel = ""
        if cfg.step_shots in ("fail", "all"):
            try:
                await page.screenshot(
                    path=str(shots_dir / f"case_{index}_step_{failed_step_no}_failure.png")
                )
            except Exception:
                pass
    finally:
        try:
            if page.video:
                temp_video_path = await page.video.path()
        except Exception:
            temp_video_path = ""
        if cfg.tracing:
            try:
                trace_file = flow_path / "traces"
                trace_file.mkdir(parents=True, exist_ok=True)
                dest = trace_file / f"case_{index}.zip"
                await context.tracing.stop(path=str(dest))
                trace_path = config.static_url(cfg.project_name, cfg.flow_name,
                                                "traces", dest.name)
            except Exception:
                trace_path = ""
        try:
            await context.close()
        except Exception:
            pass
        if temp_video_path and os.path.exists(temp_video_path):
            target_video = video_dir / f"scenario_{index}.webm"
            try:
                if target_video.exists():
                    target_video.unlink()
                shutil.move(temp_video_path, target_video)
            except Exception:
                pass

    video_url = (
        config.static_url(cfg.project_name, cfg.flow_name, "videos", f"scenario_{index}.webm")
        if temp_video_path else ""
    )

    # ── Optional visual regression ────────────────────────────────────
    visual: dict = {}
    if cfg.visual_check and screenshot_rel:
        baseline_dir = flow_path / "baseline"
        baseline_dir.mkdir(parents=True, exist_ok=True)
        baseline = baseline_dir / f"case_{index}.png"
        current = shots_dir / Path(screenshot_rel).name
        if not baseline.exists() and current.exists():
            try:
                shutil.copy(current, baseline)
                visual = {"baseline_created": True}
                await emit_log("🖼️ Visual baseline captured for future comparisons.")
            except Exception:
                pass
        elif baseline.exists() and current.exists():
            visual = compare_visual(baseline, current, flow_path / "baseline" / f"diff_{index}.png")
            if visual.get("supported") and visual.get("diff_percent", 0) > 1.0:
                await emit_log(
                    f"🖼️ Visual difference of {visual['diff_percent']}% against the baseline."
                )
            elif not visual.get("supported"):
                await emit_log(f"🖼️ {visual.get('note')}")

    # Keep what this scenario observed about the application. Written every run,
    # so the profile sharpens over time without anyone maintaining it.
    try:
        if ctx.get("observed_fields"):
            learned.remember_fields(cfg.project_name, cfg.flow_name, ctx["observed_fields"])
        for message in evidence.messages:
            if message.get("kind") in ("error", "success"):
                learned.remember_message_selector(cfg.project_name, message.get("selector", ""))
    except Exception as err:      # learning must never break a run
        log.info("Could not persist learned profile: %s", err)

    duration = int(time.time() - start_time)
    result = {
        "index": index,
        "name": name,
        "label": label,
        # The scenario this test case belongs to, and which case of it this is.
        # Both empty for a flat suite, which is what makes the old presentation
        # survive untouched.
        "parent": scenario.get("_parent_name", ""),
        "case": scenario.get("_case_name", ""),
        "case_no": scenario.get("_case_no", 0),
        "case_total": scenario.get("_case_total", 0),
        "data_row": scenario.get("_data_row") or {},
        "type": scenario.get("type", ""),
        "status": status,
        "verdict": verdict_obj.to_dict() if verdict_obj else {},
        "rule_results": [r for r in (evidence.rule_results or []) if not r["skipped"]],
        "captured": {k: v for k, v in runtime_pool.items()
                     if k not in ("run_token", "api_body") and not k.startswith("dep_")},
        "notes": evidence.notes,
        "screenshot": screenshot_rel,
        "video": video_url,
        "trace": trace_path,
        "duration": duration,
        "steps_executed": steps_executed,
        "step_durations": step_durations,
        "expected": expected,
        "messages": evidence.messages,
        "dialogs": evidence.dialogs,
        "downloads": evidence.downloads,
        "assertions": evidence.assertions,
        "healed_steps": evidence.healed_steps,
        "console_errors": console_errors,
        "network_errors": network_errors,
        "visual": visual,
        "failed_step": failed_step_no if status in ("FAIL", "CRASH") else 0,
    }
    if error_text:
        result["error"] = error_text

    state.completed += 1
    if status == verdict_engine.PASS:
        state.passed += 1
    elif status == verdict_engine.UNVERIFIED:
        state.unverified += 1
    elif status == "CRASH":
        state.crashed += 1
    elif status != "CANCELLED":
        state.failed += 1

    icon = {"PASS": "🟢", "FAIL": "🔴", "UNVERIFIED": "🟡", "CRASH": "💥",
            "CANCELLED": "⏹️"}.get(status, "⚪")
    await emit({
        "event": "test_end",
        "index": index,
        "name": name,
        "status": status,
        "verdict": result["verdict"],
        "screenshot": f"{screenshot_rel}?t={int(time.time())}" if screenshot_rel else "",
        "video": video_url,
        "trace": trace_path,
        "duration": duration,
        "console_errors": console_errors,
        "network_errors": network_errors,
        "steps_executed": steps_executed,
        "step_durations": step_durations,
        "expected": expected,
        "messages": evidence.messages,
        "downloads": evidence.downloads,
        "healed_steps": evidence.healed_steps,
        "visual": visual,
        "progress": {"completed": state.completed, "total": state.total},
        **({"error": f"{icon} {error_text}"} if error_text else {}),
    })
    if status == "CRASH":
        await emit({"event": "log", "index": index,
                    "message": f"💥 [Crash Recovery] Scenario #{index} crashed the browser. "
                               f"Continuing with the next scenario..."})
    return result


# ──────────────────────────────────────────────────────────────────────
# Suite orchestration
# ──────────────────────────────────────────────────────────────────────
def load_project_config(project_name: str) -> dict:
    path = config.project_config_file(project_name)
    if not path.exists():
        return {}
    try:
        with open(path, "r", encoding="utf-8") as fh:
            return json.load(fh)
    except Exception:
        return {}


def resolve_credentials(project_name: str, role: str = "") -> dict:
    """Credentials for a project, or for one named role within it.

    ``roles`` in the project config looks like::

        "roles": {"approver": {"username": "...", "password": "..."}}
    """
    project_cfg = load_project_config(project_name)
    if role:
        roles = project_cfg.get("roles") or {}
        chosen = dict(roles.get(role) or {})
        if chosen.get("username"):
            return chosen
        return {}
    creds = dict(project_cfg.get("credentials") or {})
    if creds.get("username"):
        return creds
    # The gitignored per-project store — where the recorder now puts the
    # literals it scrubbed out of actions.json. Checked after the project
    # config so an explicit setting still wins.
    from . import credentials as credential_store

    stored = credential_store.load_stored(project_name)
    if stored.get("username"):
        stored.update({k: v for k, v in creds.items() if v})
        return stored
    if config.GLOBAL_LOGIN_DATA_FILE.exists():
        try:
            with open(config.GLOBAL_LOGIN_DATA_FILE, "r", encoding="utf-8") as fh:
                legacy = json.load(fh)
            legacy.update({k: v for k, v in creds.items() if v})
            return legacy
        except Exception:
            pass
    return creds


def resolve_target_url(project_name: str, flow_meta: dict, env_name: str = "") -> str:
    """Apply the selected environment's base URL to the recorded target URL."""
    recorded = (flow_meta or {}).get("target_url", "") or ""
    project_cfg = load_project_config(project_name)
    environments = project_cfg.get("environments") or {}
    chosen = env_name or project_cfg.get("active_env") or ""
    base = (environments.get(chosen) or {}).get("base_url") if chosen else ""
    if not base or not recorded:
        return recorded
    try:
        from urllib.parse import urlsplit, urlunsplit

        rec = urlsplit(recorded)
        new = urlsplit(base if "//" in base else f"https://{base}")
        return urlunsplit((new.scheme or rec.scheme, new.netloc or rec.netloc,
                           rec.path, rec.query, rec.fragment))
    except Exception:
        return recorded


def select_scenarios(raw_cases: list, cfg: RunConfig) -> list[dict]:
    """Expand scenarios into test cases, then apply the filters and the limit."""
    # A scenario with a data_table is an objective plus N concrete cases. Expand
    # first so everything downstream — ids, filters, the report — counts real
    # test cases rather than templates.
    raw_cases = coverage.expand_scenarios(list(raw_cases))
    for original_index, scenario in enumerate(raw_cases, start=1):
        scenario["_original_id"] = original_index

    # Source first, and strictly: asking for Claude's cases must never quietly
    # run Gemini's because none matched.
    cases = case_sources.filter_by_sources(list(raw_cases), cfg.sources)

    if cfg.case_ids.strip():
        wanted = {c.strip() for c in cfg.case_ids.split(",") if c.strip()}
        subset = [c for c in cases if str(c["_original_id"]) in wanted]
        if subset:
            cases = subset

    if cfg.test_types.strip():
        selected = [t.strip().lower() for t in cfg.test_types.split(",") if t.strip()]
        filtered = [
            c for c in cases
            if any(st in (c.get("type", "").lower()) or (c.get("type", "").lower()) in st
                   for st in selected)
        ]
        if filtered:
            cases = filtered

    if cfg.num_cases and cfg.num_cases != -1:
        cases = cases[: cfg.num_cases]

    for run_seq, scenario in enumerate(cases, start=1):
        scenario["_run_seq"] = run_seq
    return cases


def expand_data_matrix(cases: list[dict], rows: list[dict]) -> list[dict]:
    """Duplicate every scenario once per CSV row for data-driven execution.

    A scenario that carries its own ``data_table`` is left alone. Multiplying
    the two would be a cartesian product — 9 enumerated values × 6 CSV rows —
    which is never what anyone means, and it made the global CSV useless for
    per-field permutations.
    """
    if not rows:
        return cases
    expanded: list[dict] = []
    for scenario in cases:
        if scenario.get("_parent_name") or coverage.has_data_table(scenario):
            expanded.append(scenario)
            continue
        for row_no, row in enumerate(rows, start=1):
            clone = json.loads(json.dumps(scenario))
            clone["_data_row"] = row
            clone["_label"] = f"row {row_no}"
            clone["_original_id"] = scenario["_original_id"]
            expanded.append(clone)
    return expanded


async def execute_suite(cfg: RunConfig, emit: Emit,
                        state: RunState | None = None) -> dict:
    """Run a whole suite and persist the report. Returns the report dict.

    Whatever happens, the run leaves "running": an exception marks it "error"
    here, once, rather than relying on every caller to remember (the scheduler
    did not, so a scheduled run that raised reported "running" forever).
    """
    state = state or register_run(cfg.project_name, cfg.flow_name)
    try:
        return await _execute_suite(cfg, emit, state)
    except BaseException as err:
        if state.status == "running":
            state.status = "cancelled" if isinstance(err, asyncio.CancelledError) else "error"
            state.error = state.error or str(err) or type(err).__name__
            state.finished_at = time.time()
        raise
    finally:
        persist_run_state(state)


async def _execute_suite(cfg: RunConfig, emit: Emit, state: RunState) -> dict:
    flow_path = config.flow_dir(cfg.project_name, cfg.flow_name)
    cases_path = flow_path / "test_cases.json"

    if not cases_path.exists():
        state.status = "error"
        state.error = "No test cases generated for this flow."
        await emit({"event": "error", "message": state.error})
        raise FileNotFoundError(state.error)

    await emit({"event": "start", "run_id": state.run_id,
                "message": f"Loading matrix test cases for flow '{cfg.flow_name}'..."})

    with open(cases_path, "r", encoding="utf-8") as fh:
        raw_cases = json.load(fh)

    flow_meta = {}
    meta_path = flow_path / "flow.json"
    if meta_path.exists():
        try:
            with open(meta_path, "r", encoding="utf-8") as fh:
                flow_meta = json.load(fh)
        except Exception:
            flow_meta = {}

    target_url = resolve_target_url(cfg.project_name, flow_meta, cfg.env_name)
    if not target_url:
        state.status = "error"
        state.error = "This flow has no target URL recorded."
        await emit({"event": "error", "message": state.error})
        raise ValueError(state.error)

    # ── Rules and generated scenarios ─────────────────────────────────
    active_rules = rules_engine.applicable(cfg.project_name, cfg.flow_name)         if cfg.apply_rules else []
    if active_rules:
        await emit({"event": "info",
                    "message": f"📏 {len(active_rules)} rule(s) will be enforced on every "
                               f"scenario in this run."})

    generated: list[dict] = []
    rule_scenarios = rules_engine.to_scenarios(active_rules)
    if rule_scenarios:
        generated.extend(rule_scenarios)
        await emit({"event": "info",
                    "message": f"📏 {len(rule_scenarios)} scenario(s) generated directly from "
                               f"your rules."})

    if cfg.relationship_checks:
        checks, notes = contracts.relationship_checks(
            cfg.project_name, cfg.flow_name, cfg.dependencies)
        generated.extend(checks)
        for note in notes:
            await emit({"event": "log", "message": f"🔗 {note}"})
        if checks:
            await emit({"event": "info",
                        "message": f"🔗 {len(checks)} dependency relationship check(s) added "
                                   f"(dropdown membership, inactive exclusion, delete-in-use)."})

    if generated:
        base = len(raw_cases)
        for offset, scenario in enumerate(generated, start=1):
            scenario["_original_id"] = base + offset
        raw_cases = list(raw_cases) + generated

    cases = select_scenarios(raw_cases, cfg)
    data_rows = load_data_matrix(flow_path, cfg.data_matrix)
    if data_rows:
        cases = expand_data_matrix(cases, data_rows)
        await emit({"event": "info",
                    "message": f"📊 Data-driven mode: {len(data_rows)} data row(s) × "
                               f"{len(cases) // max(len(data_rows), 1)} scenario(s)."})

    # "Screens": how many browsers run test cases at the same time. Headed runs
    # get one window per screen, tiled across the desktop (app/screens.py).
    workers = max(1, min(cfg.workers, config.MAX_PARALLEL_WORKERS, max(len(cases), 1)))
    if cfg.workers > config.MAX_PARALLEL_WORKERS:
        await emit({"event": "info",
                    "message": f"ℹ️ {cfg.workers} screens requested; the maximum is "
                               f"{config.MAX_PARALLEL_WORKERS} (QA_MAX_WORKERS in .env)."})

    state.total = len(cases)
    await emit({"event": "info", "run_id": state.run_id,
                "message": f"Prepared {len(cases)} scenarios to execute in {cfg.mode} mode"
                           f"{f' on {workers} screens at the same time' if workers > 1 else ''}."})
    if cfg.verdict_mode != "off":
        await emit({"event": "info",
                    "message": f"⚖️ Verdict engine: {cfg.verdict_mode} — expectations are "
                               f"checked against on-screen evidence, not just step success."})

    seed_pool: dict = {"run_token": unique_token()}
    if cfg.api_seed:
        try:

            seeded = await api_seed.run_seed(cfg.project_name, cfg.api_seed,
                                            emit=lambda m: emit({"event": "log", "message": m}))
            seed_pool.update(seeded)
        except Exception as err:
            await emit({"event": "log", "message": f"⚠️ API pre-seeding skipped: {err}"})

    shared = {
        "flow_path": flow_path,
        "target_url": target_url,
        "total": len(cases),
        "credentials": resolve_credentials(cfg.project_name),
        "session_file": config.project_session_file(cfg.project_name),
        "seed_pool": seed_pool,
        "rules": active_rules,
    }

    videos_dir = flow_path / "videos"
    if videos_dir.exists():
        shutil.rmtree(videos_dir, ignore_errors=True)
    videos_dir.mkdir(parents=True, exist_ok=True)

    run_results: list[dict] = []
    started = time.time()

    async with async_playwright() as playwright:
        launch_args: dict[str, Any] = {"headless": cfg.headless}
        if not cfg.headless:
            slow_mo = pacing.pace_for(cfg.speed).slow_mo_headed
            if slow_mo:
                launch_args["slow_mo"] = slow_mo
            launch_args["args"] = screens.launch_args(0 if workers > 1 else None,
                                                      workers, cfg.headless)
        browser = await playwright.chromium.launch(**launch_args)

        # ── Suite-level seeding ───────────────────────────────────────
        # Masters marked run_once in their contract are created once here and
        # their values published to every scenario, instead of being recreated
        # (and duplicated) ahead of all N scenarios.
        seed_once, per_scenario = seedable_dependencies(cfg)
        if seed_once:
            await emit({"event": "info",
                        "message": f"🌱 Seeding prerequisites once for the whole run: "
                                   f"{', '.join(seed_once)}"})
            try:
                seeded = await seed_prerequisites(browser, cfg, shared, emit, state, seed_once)
                shared["seeded_vars"] = seeded
                if seeded:
                    await emit({"event": "info",
                                "message": "🌱 Available to every scenario: " +
                                           ", ".join(f"{{{{{k}}}}}={v!r}"
                                                     for k, v in seeded.items())})
                else:
                    await emit({"event": "log",
                                "message": "🌱 Seeding produced no variables — add a "
                                           "'provides' entry to the master's contract."})
            except RunCancelled:
                raise
            except Exception as err:  # noqa: BLE001 - seeding must not kill the run
                log.exception("Seeding failed")
                await emit({"event": "error",
                            "message": f"⚠️ Seeding failed: {humanize_error(err)}. "
                                       f"Scenarios will run without the seeded values."})

        # Headed: each extra screen is its own browser window in its own tile.
        # Headless: one browser is enough; screens are just concurrent contexts.
        screen_browsers = [browser]
        if not cfg.headless and workers > 1:
            for slot in range(1, workers):
                extra = dict(launch_args)
                extra["args"] = screens.launch_args(slot, workers, cfg.headless)
                screen_browsers.append(await playwright.chromium.launch(**extra))
        if not cfg.headless:
            shared["video_sizes"] = {id(b): await measure_video_size(b)
                                     for b in screen_browsers}

        async def guarded(scenario: dict, on_browser) -> dict | None:
            if state.cancelled:
                return None
            try:
                return await run_scenario(on_browser, scenario, cfg, shared, emit, state)
            except Exception as err:  # noqa: BLE001 - never lose the whole suite
                log.exception("Scenario crashed outside the guarded block")
                await emit({"event": "log",
                            "message": f"💥 Scenario #{scenario.get('_original_id')} "
                                       f"aborted: {humanize_error(err)}"})
                return {
                    "index": scenario.get("_original_id"),
                    "name": scenario.get("scenario_name", "unknown"),
                    "parent": scenario.get("_parent_name", ""),
                    "case": scenario.get("_case_name", ""),
                    "status": "CRASH",
                    "duration": 0,
                    "steps_executed": [],
                    "step_durations": [],
                    "expected": scenario.get("expected_result", ""),
                    "error": humanize_error(err),
                    "verdict": {"status": "CRASH", "reason": humanize_error(err)},
                }

        # Every screen takes the next waiting test case, so a screen that gets
        # quick cases is not left idle while another works through slow ones.
        pending: asyncio.Queue = asyncio.Queue()
        for scenario in cases:
            pending.put_nowait(scenario)

        async def screen(slot: int) -> None:
            on_browser = screen_browsers[slot] if len(screen_browsers) > 1 else browser
            while not state.cancelled:
                try:
                    scenario = pending.get_nowait()
                except asyncio.QueueEmpty:
                    return
                result = await guarded(scenario, on_browser)
                if result:
                    run_results.append(result)

        await asyncio.gather(*(screen(slot) for slot in range(workers)))

        for each in screen_browsers:
            try:
                await each.close()
            except Exception:
                pass

    run_results.sort(key=lambda r: (r.get("index") or 0, r.get("label") or ""))
    report = build_report(cfg, run_results, started, state)
    persist_report(cfg, report, run_results)

    state.report = report
    state.finished_at = time.time()
    state.status = "cancelled" if state.cancelled else "complete"

    if state.cancelled:
        await emit({"event": "cancelled", "run_id": state.run_id,
                    "message": "⏹️ Execution stopped by the user. Partial report saved."})
    await emit({
        "event": "complete",
        "run_id": state.run_id,
        "summary": {
            "total": report["total_test_cases"],
            "passed": report["passed_count"],
            "failed": report["failed_count"],
            "unverified": report["unverified_count"],
            "crashed": report["crash_count"],
        },
        "message": f"Execution finished. Report saved to projects/{cfg.project_name}/flows/"
                   f"{cfg.flow_name}/report.json",
    })
    return report


def build_report(cfg: RunConfig, run_results: list[dict], started: float,
                 state: RunState) -> dict:
    passed = sum(1 for r in run_results if r.get("status") == verdict_engine.PASS)
    failed = sum(1 for r in run_results if r.get("status") == verdict_engine.FAIL)
    unverified = sum(1 for r in run_results if r.get("status") == verdict_engine.UNVERIFIED)
    crashed = sum(1 for r in run_results if r.get("status") == "CRASH")
    console_errors = sorted({e for r in run_results for e in r.get("console_errors", [])})
    network_errors = sorted({e for r in run_results for e in r.get("network_errors", [])})

    # Roll the per-scenario rule outcomes up into one line per rule, so the
    # report answers "did my rules hold?" without reading every scenario.
    rule_summary: dict[str, dict] = {}
    for result in run_results:
        for entry in result.get("rule_results", []):
            summary = rule_summary.setdefault(entry["id"], {
                "id": entry["id"], "title": entry["title"], "kind": entry["kind"],
                "severity": entry["severity"], "checked": 0, "violations": 0,
                "examples": [],
            })
            summary["checked"] += 1
            if not entry["ok"]:
                summary["violations"] += 1
                if len(summary["examples"]) < 3:
                    summary["examples"].append(
                        {"scenario": result.get("name"), "detail": entry["detail"]})

    return {
        "run_id": state.run_id,
        "flow_name": cfg.flow_name,
        "project_name": cfg.project_name,
        "execution_date": time.strftime("%Y-%m-%d %H:%M:%S"),
        "duration_seconds": int(time.time() - started),
        "wall_clock_seconds": int(time.time() - started),
        "test_types_run": cfg.test_types,
        "mode": cfg.mode,
        "workers": cfg.workers,
        "verdict_mode": cfg.verdict_mode,
        "environment": cfg.env_name,
        "triggered_by": cfg.triggered_by,
        "total_test_cases": len(run_results),
        "passed_count": passed,
        "failed_count": failed,
        "unverified_count": unverified,
        "crash_count": crashed,
        "cancelled": state.cancelled,
        "console_errors": console_errors,
        "network_errors": network_errors,
        # Which detection layers were active. The document that specified DOM-diff
        # asked for "a note in the report", because turning it on moves verdict
        # numbers once: a positive scenario that had been silently ignoring an
        # on-screen error now correctly fails.
        "message_detection": (["seed-selectors"]
                              + (["dom-diff"] if config.DOM_DIFF_MESSAGES else [])
                              + (["strict-confidence"] if config.STRICT_MESSAGE_CONFIDENCE
                                 else [])),
        "rules_enforced": sorted(rule_summary.values(), key=lambda r: -r["violations"]),
        "rules_violated": sum(1 for r in rule_summary.values() if r["violations"]),
        # Test cases rolled back up under the scenario they belong to. Nine data
        # rows against one field read as one finding — "7 of 9 passed" — instead
        # of nine unrelated lines that look like nine separate defects.
        "scenario_groups": coverage.group_results(run_results),
        "total_scenarios": len({r.get("parent") or r.get("name") for r in run_results}),
        "checkpoints": run_results,
    }


def persist_report(cfg: RunConfig, report: dict, run_results: list[dict]) -> None:
    """Write report.json, the timestamped copy, run history, bugs and JUnit XML."""
    flow_path = config.flow_dir(cfg.project_name, cfg.flow_name)
    reports_dir = flow_path / "reports"
    reports_dir.mkdir(parents=True, exist_ok=True)

    storage.write_json(flow_path / "report.json", report)

    run_stamp = int(time.time())
    report_file = reports_dir / f"report_{run_stamp}.json"
    if report_file.exists():
        # Two runs of one flow finishing in the same second must not overwrite
        # each other's archived report; the history row points at this name.
        suffix = report.get("run_id") or uuid.uuid4().hex[:6]
        report_file = reports_dir / f"report_{run_stamp}_{suffix}.json"
    storage.write_json(report_file, report)

    history_entry = {
        "run_id": report.get("run_id"),
        "stamp": run_stamp,
        "execution_date": report["execution_date"],
        "duration_seconds": report["duration_seconds"],
        "total": report["total_test_cases"],
        "passed": report["passed_count"],
        "failed": report["failed_count"],
        "unverified": report["unverified_count"],
        "crashed": report["crash_count"],
        "test_types": report["test_types_run"],
        "mode": report["mode"],
        "verdict_mode": report.get("verdict_mode"),
        "triggered_by": report.get("triggered_by"),
        "report_file": f"reports/{report_file.name}",
    }
    # Read-modify-write under one lock: two runs finishing together must both
    # land in the history rather than the second silently dropping the first.
    with storage.update_json(flow_path / "run_history.json", list) as doc:
        history = doc.data if isinstance(doc.data, list) else []
        history.append(history_entry)
        doc.data = history[-200:]

    # Bug list: refresh this flow's entries only.
    bugs_path = config.project_dir(cfg.project_name) / "bugs.json"
    new_bugs: list = []

    for result in run_results:
        if result.get("status") not in (verdict_engine.FAIL, "CRASH", verdict_engine.UNVERIFIED):
            continue
        steps = result.get("steps_executed", [])
        failed_step_no = result.get("failed_step") or len(steps)
        failed_step_desc = steps[failed_step_no - 1] if 0 < failed_step_no <= len(steps) else (
            steps[-1] if steps else "None")
        verdict_info = result.get("verdict") or {}
        reason = result.get("error") or verdict_info.get("reason", "Assertion failed")
        severity = {
            "CRASH": "high", verdict_engine.FAIL: "high", verdict_engine.UNVERIFIED: "low",
        }.get(result.get("status"), "medium")

        new_bugs.append({
            "flow_name": cfg.flow_name,
            "scenario_name": result.get("name"),
            "scenario_type": result.get("type", ""),
            "status": result.get("status"),
            "severity": severity,
            "failed_step_num": failed_step_no,
            "failed_step_desc": failed_step_desc,
            "error_message": reason,
            "expected_result": result.get("expected", ""),
            "observed_messages": [m.get("text") for m in result.get("messages", [])][:5],
            "description": (
                f"{cfg.flow_name} - {result.get('name')}: {result.get('status')} — {reason}"
            ),
            "screenshot": result.get("screenshot", ""),
            "video": result.get("video", ""),
            "run_id": report.get("run_id"),
            "timestamp": report["execution_date"],
        })

    # bugs.json is project-wide, so runs of *different* flows race on it too.
    with storage.update_json(bugs_path, list) as doc:
        bugs = doc.data if isinstance(doc.data, list) else []
        doc.data = [b for b in bugs if b.get("flow_name") != cfg.flow_name] + new_bugs

    try:
        junit = build_junit_xml(report)
        storage.write_text(reports_dir / f"junit_{run_stamp}.xml", junit)
        storage.write_text(flow_path / "junit.xml", junit)
    except Exception as err:
        log.warning("Could not write JUnit XML: %s", err)

    try:
        from . import db

        db.record_run(report)
    except Exception as err:
        log.warning("Could not write run to the database: %s", err)


def _xml_escape(text: str) -> str:
    return (
        str(text)
        .replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        .replace('"', "&quot;").replace("'", "&apos;")
    )


def build_junit_xml(report: dict) -> str:
    """JUnit XML so Jenkins / GitHub Actions / Azure DevOps can read the results."""
    cases = report.get("checkpoints", [])
    failures = report.get("failed_count", 0) + report.get("crash_count", 0)
    skipped = report.get("unverified_count", 0)
    lines = [
        '<?xml version="1.0" encoding="UTF-8"?>',
        f'<testsuites name="{_xml_escape(report.get("project_name", ""))}" '
        f'tests="{len(cases)}" failures="{failures}" skipped="{skipped}" '
        f'time="{report.get("duration_seconds", 0)}">',
        f'  <testsuite name="{_xml_escape(report.get("flow_name", ""))}" '
        f'tests="{len(cases)}" failures="{failures}" skipped="{skipped}" '
        f'time="{report.get("duration_seconds", 0)}" '
        f'timestamp="{_xml_escape(report.get("execution_date", ""))}">',
    ]
    for case in cases:
        name = _xml_escape(case.get("name", "scenario"))
        classname = _xml_escape(f"{report.get('project_name')}.{report.get('flow_name')}")
        lines.append(
            f'    <testcase classname="{classname}" name="{name}" '
            f'time="{case.get("duration", 0)}">'
        )
        status = case.get("status")
        reason = _xml_escape((case.get("verdict") or {}).get("reason")
                             or case.get("error") or "")
        if status == verdict_engine.FAIL:
            lines.append(f'      <failure message="{reason}">{reason}</failure>')
        elif status == "CRASH":
            lines.append(f'      <error message="{reason}">{reason}</error>')
        elif status == verdict_engine.UNVERIFIED:
            lines.append(f'      <skipped message="{reason}"/>')
        lines.append("    </testcase>")
    lines += ["  </testsuite>", "</testsuites>"]
    return "\n".join(lines)
